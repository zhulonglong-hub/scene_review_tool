from pathlib import Path

import pytest
from PIL import Image

from scene_review_tool.app import (
    Review,
    ReviewDatabase,
    QUALITY_NAMES,
    Sample,
    Taxonomy,
    build_export_plan,
    compose_preview,
    compose_preview_layers,
    export_review_items,
    export_quality_items,
    infer_mask_schema,
    inspect_sample_pairs,
    read_quality_progress_xlsx,
    resolve_export_start_directory,
    scan_dataset,
    validate_mask_schema,
)


def test_taxonomy_loads_project_taxonomy(tmp_path):
    taxonomy_path = tmp_path / "taxonomy.json"
    taxonomy_path.write_text(
        """
        {
          "version": "1.0",
          "domains": [
            {
              "id": "G1",
              "name_zh": "水系与湿地",
              "name_en": "Water",
              "scenes": [{"scene": "river", "elements": ["water"]}]
            }
          ]
        }
        """,
        encoding="utf-8",
    )
    taxonomy = Taxonomy(taxonomy_path)
    assert taxonomy.version == "1.0"
    assert taxonomy.find_scene("river") is not None


def test_taxonomy_can_be_omitted():
    taxonomy = Taxonomy()

    assert taxonomy.path is None
    assert taxonomy.scenes == []
    assert taxonomy.domains == []


def test_export_start_directory_falls_back_and_prefers_existing_paths(tmp_path):
    fallback = tmp_path / "fallback"
    workspace = tmp_path / "workspace"
    previous = tmp_path / "previous-export"
    fallback.mkdir()
    workspace.mkdir()

    assert resolve_export_start_directory(None, None, "", fallback) == str(fallback)
    assert resolve_export_start_directory(None, workspace, "", fallback) == str(workspace)

    previous.mkdir()
    assert resolve_export_start_directory(previous, workspace, "", fallback) == str(previous)


def test_mirrored_subdirectories_scan(tmp_path):
    dataset = tmp_path / "dataset"
    image_dir = dataset / "group1" / "images"
    mask_dir = dataset / "group1" / "labels"
    image_dir.mkdir(parents=True)
    mask_dir.mkdir(parents=True)
    Image.new("RGB", (8, 8), (10, 20, 30)).save(image_dir / "a.png")
    Image.new("L", (8, 8), 255).save(mask_dir / "a.png")

    samples, warnings = scan_dataset("demo", dataset, dataset, "mirrored_subdirectories", "images", "labels")

    assert len(samples) == 1
    assert warnings == []
    assert samples[0].parent_group_id == "group1"


def test_voc_segmentation_scan_uses_split_file(tmp_path):
    dataset = tmp_path / "voc"
    image_dir = dataset / "JPEGImages"
    mask_dir = dataset / "SegmentationClass"
    split_dir = dataset / "ImageSets" / "Segmentation"
    image_dir.mkdir(parents=True)
    mask_dir.mkdir(parents=True)
    split_dir.mkdir(parents=True)
    for name in ("a", "b"):
        Image.new("RGB", (8, 8), (10, 20, 30)).save(image_dir / f"{name}.jpg")
        Image.new("L", (8, 8), 1).save(mask_dir / f"{name}.png")
    (split_dir / "test.txt").write_text("b\n", encoding="utf-8")

    samples, warnings = scan_dataset(
        "voc-demo", dataset, dataset, "voc_segmentation", "JPEGImages", "SegmentationClass", "test"
    )

    assert warnings == []
    assert len(samples) == 1
    assert Path(samples[0].image_path).stem == "b"
    assert Path(samples[0].mask_path).stem == "b"
    assert samples[0].parent_group_id == "test"


def test_generic_relative_path_scan_pairs_duplicate_stems_safely(tmp_path):
    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    for group in ("north", "south"):
        (image_root / group).mkdir(parents=True)
        (mask_root / group).mkdir(parents=True)
        Image.new("RGB", (8, 8), (10, 20, 30)).save(image_root / group / "tile.jpg")
        Image.new("L", (8, 8), 1).save(mask_root / group / "tile.png")

    samples, warnings = scan_dataset(
        "relative-demo", image_root, mask_root, "relative_path_stem", "", "", "all", True
    )

    assert len(samples) == 2
    assert warnings == []
    assert {Path(sample.image_path).parent.name for sample in samples} == {"north", "south"}


def test_generic_scan_pairs_mask_with_filename_suffix(tmp_path):
    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    image_root.mkdir()
    mask_root.mkdir()
    Image.new("RGB", (8, 8), (10, 20, 30)).save(image_root / "P0018.png")
    Image.new("L", (8, 8), 1).save(mask_root / "P0018_instance_color_RGB.png")

    samples, warnings = scan_dataset(
        "suffix-demo",
        image_root,
        mask_root,
        "same_stem",
        "",
        "",
        "all",
        True,
        "",
        "_instance_color_RGB",
    )

    assert warnings == []
    assert len(samples) == 1
    assert Path(samples[0].image_path).name == "P0018.png"
    assert Path(samples[0].mask_path).name == "P0018_instance_color_RGB.png"


def test_generic_scan_pairs_mask_with_filename_prefix_and_suffix(tmp_path):
    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    image_root.mkdir()
    mask_root.mkdir()
    Image.new("RGB", (8, 8), (10, 20, 30)).save(image_root / "tile.png")
    Image.new("L", (8, 8), 1).save(mask_root / "mask_tile_label.png")

    samples, warnings = scan_dataset(
        "prefix-suffix-demo",
        image_root,
        mask_root,
        "same_stem",
        "",
        "",
        "all",
        True,
        "mask_",
        "_label",
    )

    assert warnings == []
    assert len(samples) == 1
    assert Path(samples[0].mask_path).name == "mask_tile_label.png"


def test_generic_same_stem_does_not_guess_ambiguous_pairs(tmp_path):
    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    for group in ("north", "south"):
        (image_root / group).mkdir(parents=True)
        (mask_root / group).mkdir(parents=True)
        Image.new("RGB", (8, 8), (10, 20, 30)).save(image_root / group / "tile.jpg")
        Image.new("L", (8, 8), 1).save(mask_root / group / "tile.png")

    samples, warnings = scan_dataset(
        "ambiguous-demo", image_root, mask_root, "same_stem", "", "", "all", True
    )

    assert samples == []
    assert warnings == ["ambiguous duplicate pairing key: tile"]


def test_deepglobe_suffix_scan_pairs_shared_folder_and_reports_missing_counterparts(tmp_path):
    dataset = tmp_path / "deepglobe-road"
    dataset.mkdir()
    for sample_id in ("100034", "100081"):
        Image.new("RGB", (8, 8), (10, 20, 30)).save(dataset / f"{sample_id}_sat.jpg")
        Image.new("RGB", (8, 8), (255, 255, 255)).save(dataset / f"{sample_id}_mask.png")
    Image.new("RGB", (8, 8), (10, 20, 30)).save(dataset / "image_only_sat.jpg")
    Image.new("RGB", (8, 8), (255, 255, 255)).save(dataset / "mask_only_mask.png")
    (dataset / "100034_mask.png.aux.xml").write_text("metadata", encoding="utf-8")

    samples, warnings = scan_dataset(
        "deepglobe-road", dataset, dataset, "deepglobe_suffix", "", "", "all", True
    )

    assert len(samples) == 2
    assert {
        (Path(sample.image_path).name, Path(sample.mask_path).name)
        for sample in samples
    } == {
        ("100034_sat.jpg", "100034_mask.png"),
        ("100081_sat.jpg", "100081_mask.png"),
    }
    assert warnings == [
        "missing DeepGlobe mask for id: image_only",
        "missing DeepGlobe image for id: mask_only",
    ]


def test_binary_mask_value_one_is_rendered_as_foreground(tmp_path):
    image_path = tmp_path / "image.png"
    mask_path = tmp_path / "mask.png"
    Image.new("RGB", (2, 1), (10, 20, 30)).save(image_path)
    mask = Image.new("L", (2, 1), 0)
    mask.putpixel((1, 0), 1)
    mask.save(mask_path)

    preview = compose_preview(str(image_path), str(mask_path), True, 50, {1}).convert("RGB")

    assert preview.getpixel((0, 0)) == (10, 20, 30)
    assert preview.getpixel((1, 0)) != (10, 20, 30)


def test_grayscale_multiclass_mask_is_detected_and_background_is_suggested(tmp_path):
    image_path = tmp_path / "image.png"
    mask_path = tmp_path / "mask.png"
    Image.new("RGB", (3, 1), (10, 20, 30)).save(image_path)
    mask = Image.new("L", (3, 1))
    mask.putdata([0, 1, 2])
    mask.save(mask_path)
    sample = Sample("gray", str(image_path), str(mask_path), "g")

    inspection = inspect_sample_pairs([sample])
    schema = infer_mask_schema(inspection)

    assert inspection["mask_encoding"] == "indexed"
    assert {label["key"] for label in inspection["mask_labels"]} == {"i:0", "i:1", "i:2"}
    assert next(label for label in schema["labels"] if label["key"] == "i:0")["role"] == "background"
    assert {label["name"] for label in schema["labels"] if label["role"] == "class"} == {
        "class_1",
        "class_2",
    }


def test_palette_mask_preserves_palette_indices(tmp_path):
    image_path = tmp_path / "image.png"
    mask_path = tmp_path / "palette.png"
    Image.new("RGB", (2, 1), (10, 20, 30)).save(image_path)
    mask = Image.new("P", (2, 1))
    palette = [0] * 768
    palette[3:6] = [0, 255, 0]
    mask.putpalette(palette)
    mask.putdata([0, 1])
    mask.save(mask_path)

    inspection = inspect_sample_pairs([Sample("palette", str(image_path), str(mask_path), "g")])

    assert inspection["mask_encoding"] == "indexed"
    assert {label["key"] for label in inspection["mask_labels"]} == {"i:0", "i:1"}
    assert next(label for label in inspection["mask_labels"] if label["key"] == "i:1")["color"] == "#00ff00"


def test_rgb_mask_keeps_full_color_labels_and_renders_each_class(tmp_path):
    image_path = tmp_path / "image.png"
    mask_path = tmp_path / "rgb-mask.png"
    Image.new("RGB", (3, 1), (20, 20, 20)).save(image_path)
    mask = Image.new("RGB", (3, 1))
    mask.putdata([(0, 0, 0), (0, 255, 0), (0, 0, 255)])
    mask.save(mask_path)
    sample = Sample("rgb", str(image_path), str(mask_path), "g")

    inspection = inspect_sample_pairs([sample])
    schema = infer_mask_schema(inspection)
    preview = compose_preview(str(image_path), str(mask_path), True, 80, None, schema).convert("RGB")

    assert inspection["mask_encoding"] == "rgb"
    assert {label["key"] for label in inspection["mask_labels"]} == {
        "rgb:0,0,0",
        "rgb:0,255,0",
        "rgb:0,0,255",
    }
    assert preview.getpixel((0, 0)) == (20, 20, 20)
    assert preview.getpixel((1, 0)) != preview.getpixel((2, 0))


def test_rgb_full_classification_mask_does_not_require_background(tmp_path):
    image_path = tmp_path / "image.png"
    mask_path = tmp_path / "rgb-mask.png"
    Image.new("RGB", (2, 1), (20, 20, 20)).save(image_path)
    mask = Image.new("RGB", (2, 1))
    mask.putdata([(0, 0, 0), (255, 0, 0)])
    mask.save(mask_path)
    schema = {
        "schema_version": 3,
        "encoding": "rgb",
        "background_mode": "none",
        "labels": [
            {"key": "rgb:0,0,0", "value": [0, 0, 0], "role": "class", "name": "water", "color": "#0000ff"},
            {"key": "rgb:255,0,0", "value": [255, 0, 0], "role": "class", "name": "building", "color": "#ff0000"},
        ],
    }

    validate_mask_schema(schema)
    _image, overlay = compose_preview_layers(str(image_path), str(mask_path), None, schema)

    assert overlay.getpixel((0, 0))[3] == 255
    assert overlay.getpixel((1, 0))[3] == 255
    assert overlay.getpixel((0, 0)) != overlay.getpixel((1, 0))


def test_mask_schema_requires_at_least_one_effective_class():
    schema = {
        "schema_version": 3,
        "encoding": "indexed",
        "background_mode": "explicit",
        "labels": [
            {"key": "i:0", "value": 0, "role": "background", "name": "background", "color": "#000000"}
        ],
    }

    with pytest.raises(ValueError, match="至少指定一个有效类别"):
        validate_mask_schema(schema)


def test_database_persists_samples(tmp_path):
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    taxonomy = Taxonomy()
    dataset_id = db.create_dataset("demo", tmp_path, tmp_path, {"pairing": {}}, taxonomy)
    assert dataset_id == 1
    db.close()


def test_common_scene_count_is_assigned_sample_count(tmp_path):
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    taxonomy = Taxonomy()
    dataset_id = db.create_dataset("demo", tmp_path, tmp_path, {"pairing": {}}, taxonomy)
    sample = Sample("s1", "image.png", "mask.png", "group")
    db.add_samples([sample], dataset_id)
    review = Review(
        quality_status="accepted",
        scene_status="assigned",
        level1_id="G1",
        level1_name="水系与湿地",
        primary_level2_scene="river",
        scene_source="taxonomy",
    )

    db.save_review(sample.id, dataset_id, review)
    db.save_review(sample.id, dataset_id, review)

    common = db.common_scenes(dataset_id)
    assert len(common) == 1
    assert common[0]["scene_name_en"] == "river"
    assert common[0]["assigned_count"] == 1
    db.close()


def test_export_plan_separates_quality_and_scene_exports():
    items = [
        (Sample("a", "a.png", "a-mask.png", "g"), Review(quality_status="accepted")),
        (
            Sample("b", "b.png", "b-mask.png", "g"),
            Review(quality_status="accepted", scene_status="assigned", primary_level2_scene="river"),
        ),
        (
            Sample("c", "c.png", "c-mask.png", "g"),
            Review(quality_status="accepted", scene_status="uncertain", primary_level2_scene="lake"),
        ),
        (
            Sample("d", "d.png", "d-mask.png", "g"),
            Review(quality_status="rejected", scene_status="assigned", primary_level2_scene="river"),
        ),
        (
            Sample("e", "e.png", "e-mask.png", "g"),
            Review(quality_status="accepted", scene_status="assigned"),
        ),
    ]

    quality_plan = build_export_plan(items, group_by_scene=False)
    scene_plan = build_export_plan(items, group_by_scene=True)

    assert [sample.id for sample, _ in quality_plan.items] == ["a", "b", "c", "e"]
    assert [sample.id for sample, _ in scene_plan.items] == ["b"]
    assert scene_plan.accepted_total == 4
    assert scene_plan.assigned_total == 1
    assert scene_plan.unassigned_total == 1
    assert scene_plan.uncertain_total == 1
    assert scene_plan.invalid_assigned_total == 1


def test_export_review_items_writes_status_manifest_and_expected_layout(tmp_path):
    image_path = tmp_path / "source.jpg"
    mask_path = tmp_path / "source.png"
    Image.new("RGB", (2, 2), (10, 20, 30)).save(image_path)
    Image.new("L", (2, 2), 1).save(mask_path)
    sample = Sample("sample-id", str(image_path), str(mask_path), "g")
    review = Review(quality_status="accepted", scene_status="assigned", primary_level2_scene="river")
    plan = build_export_plan([(sample, review)], group_by_scene=True)
    root = tmp_path / "export"

    exported = export_review_items(plan, root, group_by_scene=True)

    assert exported == 1
    assert (root / "river" / "images" / "source.jpg").is_file()
    assert (root / "river" / "masks" / "source.png").is_file()
    manifest = (root / "export_manifest.csv").read_text(encoding="utf-8-sig")
    assert "quality_status,scene_status,scene" in manifest
    assert "accepted,assigned,river" in manifest


def test_quality_export_copies_accepted_sample_without_scene(tmp_path):
    image_path = tmp_path / "quality-source.jpg"
    mask_path = tmp_path / "quality-source.png"
    Image.new("RGB", (2, 2), (10, 20, 30)).save(image_path)
    Image.new("L", (2, 2), 1).save(mask_path)
    sample = Sample("quality-id", str(image_path), str(mask_path), "g")
    plan = build_export_plan([(sample, Review(quality_status="accepted"))], group_by_scene=False)
    root = tmp_path / "quality-export"

    exported = export_review_items(plan, root, group_by_scene=False)

    assert exported == 1
    assert (root / "images" / "quality-source.jpg").is_file()
    assert (root / "masks" / "quality-source.png").is_file()


def test_quality_export_preserves_names_classes_and_literal_notes(tmp_path):
    from openpyxl import load_workbook

    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    items = []
    for i, status in enumerate(("accepted", "rejected", "needs_correction", "unreviewed")):
        group = str(i)
        image = image_root / group / "P0018.png"
        mask = mask_root / group / "P0018_instance_color_RGB.png"
        image.parent.mkdir(parents=True)
        mask.parent.mkdir(parents=True)
        Image.new("RGB", (2, 2), (i, 20, 30)).save(image)
        Image.new("RGB", (2, 2), (40, i, 60)).save(mask)
        items.append((Sample(group, str(image), str(mask), group),
                      Review(quality_status=status, reviewer_note="=文字说明\n第二行,备注")))
    root = tmp_path / "export"
    assert export_quality_items(items, root, image_root, mask_root,
                                {"accepted", "rejected", "needs_correction"}) == 3
    for i, name in enumerate(("合格", "不合格", "需修改")):
        assert (root / name / "images" / str(i) / "P0018.png").is_file()
        assert (root / name / "masks" / str(i) / "P0018_instance_color_RGB.png").is_file()
    workbook = load_workbook(root / "质量审核记录.xlsx")
    sheet = workbook.active
    assert sheet.max_row == 4
    assert [sheet.cell(i, 4).value for i in range(2, 5)] == ["合格", "不合格", "需修改"]
    assert sheet["E2"].value == "=文字说明\n第二行,备注"
    assert sheet["E2"].data_type == "s"
    workbook.close()


def test_sample_list_displays_sequence_numbers(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    from scene_review_tool.app import MainWindow

    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset = db.create_dataset("demo", tmp_path, tmp_path, {}, Taxonomy())
    samples = [
        Sample("a", str(tmp_path / "P0018.png"), str(tmp_path / "P0018_mask.png"), ""),
        Sample("b", str(tmp_path / "P0019.png"), str(tmp_path / "P0019_mask.png"), ""),
    ]
    db.add_samples(samples, dataset)
    window.db = db
    window.dataset_id = dataset
    monkeypatch.setattr(window, "refresh_image", lambda: None)

    window.reload_samples()

    assert window.sample_list.item(0).text().startswith("1. · P0018.png")
    assert window.sample_list.item(1).text().startswith("2. · P0019.png")
    window.update_sample_list_item(1, samples[1], Review(quality_status="accepted"))
    assert window.sample_list.item(1).text().startswith("2. ✓ P0019.png")
    window.close()
    db.close()


def test_selected_export_exports_exactly_highlighted_samples(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox
    from scene_review_tool.app import MainWindow

    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    image_root.mkdir()
    mask_root.mkdir()
    samples = []
    for index in range(3):
        image_path = image_root / f"sample_{index}.jpg"
        mask_path = mask_root / f"sample_{index}.png"
        Image.new("RGB", (2, 2), (index, 20, 30)).save(image_path)
        Image.new("L", (2, 2), index).save(mask_path)
        samples.append(Sample(str(index), str(image_path), str(mask_path), ""))

    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset_id = db.create_dataset("demo", image_root, mask_root, {}, Taxonomy())
    db.add_samples(samples, dataset_id)
    db.save_review("0", dataset_id, Review(quality_status="accepted", reviewer_note="保留"))
    db.save_review("1", dataset_id, Review(quality_status="accepted", reviewer_note="不应导出"))
    db.save_review("2", dataset_id, Review(quality_status="rejected", reviewer_note="保留"))
    window.db = db
    window.dataset_id = dataset_id
    monkeypatch.setattr(window, "refresh_image", lambda: None)
    window.reload_samples()
    window.sample_list.clearSelection()
    window.sample_list.item(0).setSelected(True)
    window.sample_list.item(2).setSelected(True)
    assert window.selected_export_btn.text() == "导出选中（2）"

    destination = tmp_path / "exports"
    destination.mkdir()
    monkeypatch.setattr(window, "choose_export_directory", lambda _title: destination)
    monkeypatch.setattr(
        window,
        "choose_quality_export_options",
        lambda items, title, scope_text="": ({"accepted", "rejected"}, True),
    )
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
    window.export_selected()

    export_roots = list(destination.glob("selected_quality_export_*"))
    assert len(export_roots) == 1
    export_root = export_roots[0]
    assert (export_root / "合格" / "images" / "sample_0.jpg").is_file()
    assert (export_root / "合格" / "masks" / "sample_0.png").is_file()
    assert (export_root / "不合格" / "images" / "sample_2.jpg").is_file()
    assert (export_root / "不合格" / "masks" / "sample_2.png").is_file()
    assert not any(path.name == "sample_1.jpg" for path in export_root.rglob("*.jpg"))
    assert (export_root / "质量审核记录.xlsx").is_file()
    window.close()
    db.close()


def test_quality_export_report_only_and_collision_preflight(tmp_path):
    from openpyxl import load_workbook

    sample = Sample("a", str(tmp_path / "a.png"), str(tmp_path / "a_mask.png"), "")
    item = (sample, Review(quality_status="rejected", reviewer_note="原因"))
    root = tmp_path / "report"
    assert export_quality_items([item], root, tmp_path, tmp_path, {"rejected"}, False) == 1
    assert not (root / "不合格").exists()
    workbook = load_workbook(root / "质量审核记录.xlsx")
    assert workbook.active["I2"].value is None
    workbook.close()
    conflict = tmp_path / "conflict"
    with pytest.raises(ValueError, match="冲突"):
        export_quality_items([item, item], conflict, tmp_path, tmp_path, {"rejected"}, False)
    assert not conflict.exists()


def test_quality_export_independent_table_and_category_folders(tmp_path):
    from openpyxl import load_workbook

    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    image_root.mkdir()
    mask_root.mkdir()
    items = []
    for status in ("accepted", "rejected", "needs_correction"):
        image = image_root / f"{status}.png"
        mask = mask_root / f"{status}.png"
        Image.new("RGB", (2, 2), (10, 20, 30)).save(image)
        Image.new("L", (2, 2), 1).save(mask)
        items.append((Sample(status, str(image), str(mask), ""), Review(quality_status=status)))

    table_only = tmp_path / "table-only"
    assert export_quality_items(
        items, table_only, image_root, mask_root, set(),
        include_table=True, table_statuses=set(QUALITY_NAMES),
    ) == 3
    assert not (table_only / "合格").exists()
    workbook = load_workbook(table_only / "质量审核记录.xlsx")
    assert workbook.active.max_row == 4
    workbook.close()

    folders_only = tmp_path / "folders-only"
    assert export_quality_items(
        items, folders_only, image_root, mask_root, {"accepted", "needs_correction"},
        include_table=False,
    ) == 2
    assert (folders_only / "合格" / "images" / "accepted.png").is_file()
    assert (folders_only / "需修改" / "masks" / "needs_correction.png").is_file()
    assert not (folders_only / "不合格").exists()
    assert not (folders_only / "质量审核记录.xlsx").exists()

    mixed = tmp_path / "mixed"
    assert export_quality_items(
        items, mixed, image_root, mask_root, {"rejected"},
        include_table=True, table_statuses=set(QUALITY_NAMES),
    ) == 3
    assert (mixed / "不合格" / "images" / "rejected.png").is_file()
    assert not (mixed / "合格").exists()
    workbook = load_workbook(mixed / "质量审核记录.xlsx")
    assert workbook.active.max_row == 4
    assert workbook.active["I2"].value is None
    assert Path(workbook.active["I3"].value) == Path("不合格") / "images" / "rejected.png"
    workbook.close()


def test_quality_note_saves_on_selection_and_keeps_saved_scene(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    from scene_review_tool.app import MainWindow

    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset = db.create_dataset("demo", tmp_path, tmp_path, {}, Taxonomy())
    samples = [Sample(str(i), str(tmp_path / f"{i}.png"), str(tmp_path / f"{i}_mask.png"), "") for i in range(2)]
    db.add_samples(samples, dataset)
    saved = Review(quality_status="rejected", scene_status="assigned", primary_level2_scene="river")
    db.save_review("0", dataset, saved)
    window.db = db
    window.dataset_id = dataset
    window.items = db.samples(dataset)
    monkeypatch.setattr(window, "refresh_image", lambda: None)
    window.on_sample_selected(0)
    window.custom_en_edit.setText("unconfirmed candidate")
    window.note_edit.setPlainText("边界不正确\n需要检查")
    window.on_sample_selected(1)
    review = db.samples(dataset)[0][1]
    assert review.reviewer_note == "边界不正确\n需要检查"
    assert review.primary_level2_scene == "river"
    assert review.scene_status == "assigned"
    assert window.note_edit.toPlainText() == ""
    window.close()
    db.close()


def test_next_unreviewed_skips_quality_reviewed_samples_without_scenes(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox
    from scene_review_tool.app import MainWindow

    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset = db.create_dataset("demo", tmp_path, tmp_path, {}, Taxonomy())
    samples = [
        Sample(str(index), str(tmp_path / f"{index}.png"), str(tmp_path / f"{index}_mask.png"), "")
        for index in range(4)
    ]
    db.add_samples(samples, dataset)
    db.save_review("0", dataset, Review(quality_status="accepted"))
    db.save_review("1", dataset, Review(quality_status="rejected"))
    db.save_review("2", dataset, Review(quality_status="needs_correction"))
    window.db = db
    window.dataset_id = dataset
    monkeypatch.setattr(window, "refresh_image", lambda: None)
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
    window.reload_samples()

    assert window.current_index == 0
    assert window.items[1][1].scene_status == "unassigned"
    window.next_unreviewed()
    assert window.current_index == 3
    assert window.sample_list.currentRow() == 3

    window.next_unreviewed()
    assert window.current_index == 3
    window.close()
    db.close()


def test_quality_progress_excel_matches_relative_paths_and_classifies_conflicts(tmp_path):
    from openpyxl import Workbook

    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    samples = []
    for group in ("north", "south"):
        image = image_root / group / "P0018.png"
        mask = mask_root / group / "P0018_mask.png"
        image.parent.mkdir(parents=True)
        mask.parent.mkdir(parents=True)
        Image.new("RGB", (2, 2), (10, 20, 30)).save(image)
        Image.new("L", (2, 2), 1).save(mask)
        samples.append(Sample(group, str(image), str(mask), group))

    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset = db.create_dataset("demo", image_root, mask_root, {}, Taxonomy())
    db.add_samples(samples, dataset)
    db.save_review(
        "south",
        dataset,
        Review(
            quality_status="accepted",
            scene_status="assigned",
            primary_level2_scene="river",
            reviewer_note="本地备注",
        ),
    )

    excel = tmp_path / "member.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "质量审核记录"
    sheet.append([
        "样本ID", "原图名称", "掩膜名称", "质量类型", "文本描述信息备注",
        "场景名称", "原图来源路径", "掩膜来源路径", "原图导出路径", "掩膜导出路径",
    ])
    sheet.append([
        "member-north", "P0018.png", "P0018_mask.png", "合格", "成员备注", "",
        "E:/member/images/north/P0018.png", "E:/member/masks/north/P0018_mask.png", "", "",
    ])
    sheet.append([
        "member-south", "P0018.png", "P0018_mask.png", "不合格", "成员冲突", "",
        "E:/member/images/south/P0018.png", "E:/member/masks/south/P0018_mask.png", "", "",
    ])
    workbook.save(excel)
    workbook.close()

    plan = read_quality_progress_xlsx(
        excel, db.samples(dataset), image_root, mask_root
    )

    assert [entry.matched_sample_id for entry in plan.entries] == ["north", "south"]
    assert [entry.match_method for entry in plan.entries] == ["相对路径", "相对路径"]
    assert [entry.category for entry in plan.entries] == ["ready", "conflict"]

    _import_id, applied = db.apply_quality_import(dataset, plan, "fill_unreviewed")
    assert applied == 1
    reviews = {sample.id: review for sample, review in db.samples(dataset)}
    assert reviews["north"].quality_status == "accepted"
    assert reviews["north"].reviewer_note == "成员备注"
    assert reviews["south"].quality_status == "accepted"
    assert reviews["south"].reviewer_note == "本地备注"
    assert reviews["south"].primary_level2_scene == "river"
    assert db.quality_import_seen(dataset, plan.source_sha256)
    db.close()


def test_quality_progress_import_can_override_quality_without_changing_scene(tmp_path):
    from openpyxl import Workbook

    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    image_root.mkdir()
    mask_root.mkdir()
    image = image_root / "unique.png"
    mask = mask_root / "unique_mask.png"
    Image.new("RGB", (2, 2), (10, 20, 30)).save(image)
    Image.new("L", (2, 2), 1).save(mask)
    sample = Sample("local-id", str(image), str(mask), "")
    db = ReviewDatabase(tmp_path / "review.sqlite3")
    dataset = db.create_dataset("demo", image_root, mask_root, {}, Taxonomy())
    db.add_samples([sample], dataset)
    db.save_review(
        sample.id,
        dataset,
        Review(
            quality_status="accepted",
            scene_status="assigned",
            primary_level2_scene="wetland",
            reviewer_note="本地",
        ),
    )

    excel = tmp_path / "member.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "质量审核记录"
    sheet.append(["原图名称", "掩膜名称", "质量类型", "文本描述信息备注"])
    sheet.append(["unique.png", "unique_mask.png", "需修改", "成员"])
    workbook.save(excel)
    workbook.close()

    plan = read_quality_progress_xlsx(
        excel, db.samples(dataset), image_root, mask_root
    )
    assert plan.entries[0].category == "conflict"
    assert plan.entries[0].match_method == "原图名 + 掩膜名"

    _import_id, applied = db.apply_quality_import(
        dataset,
        plan,
        "fill_unreviewed",
        {"local-id": "use_imported"},
    )
    assert applied == 1
    review = db.samples(dataset)[0][1]
    assert review.quality_status == "needs_correction"
    assert review.reviewer_note == "成员"
    assert review.scene_status == "assigned"
    assert review.primary_level2_scene == "wetland"
    db.close()


def test_quality_progress_excel_rejects_ambiguous_and_internal_conflicts(tmp_path):
    from openpyxl import Workbook

    image_root = tmp_path / "images"
    mask_root = tmp_path / "masks"
    samples = []
    for group in ("a", "b"):
        image = image_root / group / "same.png"
        mask = mask_root / group / "same_mask.png"
        image.parent.mkdir(parents=True)
        mask.parent.mkdir(parents=True)
        Image.new("RGB", (2, 2), (10, 20, 30)).save(image)
        Image.new("L", (2, 2), 1).save(mask)
        samples.append(Sample(group, str(image), str(mask), group))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "质量审核记录"
    sheet.append(["样本ID", "原图名称", "掩膜名称", "质量类型", "文本描述信息备注"])
    sheet.append(["", "same.png", "same_mask.png", "合格", "无法确定目录"])
    sheet.append(["a", "same.png", "same_mask.png", "合格", "第一条"])
    sheet.append(["a", "same.png", "same_mask.png", "不合格", "第二条"])
    excel = tmp_path / "member.xlsx"
    workbook.save(excel)
    workbook.close()

    plan = read_quality_progress_xlsx(
        excel,
        [(samples[0], Review()), (samples[1], Review())],
        image_root,
        mask_root,
    )

    assert plan.entries[0].category == "unmatched"
    assert plan.entries[1].category == "duplicate_conflict"
    assert plan.entries[2].category == "duplicate_conflict"
