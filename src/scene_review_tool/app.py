from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import sqlite3
import sys
import time
from collections import Counter
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image
from PySide6.QtCore import QObject, QRunnable, QThreadPool, QTimer, Qt, Signal, Slot
from PySide6.QtGui import QColor, QImage, QPainter, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QDialog,
    QDialogButtonBox,
    QFormLayout,
    QFrame,
    QGraphicsScene,
    QGraphicsView,
    QGridLayout,
    QGroupBox,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QSlider,
    QSplitter,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
MASK_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
DEEPGLOBE_IMAGE_SUFFIX = "_sat"
DEEPGLOBE_MASK_SUFFIX = "_mask"
QUALITY_STATUSES = ["unreviewed", "accepted", "needs_correction", "rejected"]
SCENE_SOURCES = ["taxonomy", "dataset_custom", "candidate_new_taxonomy"]


@dataclass(frozen=True)
class Scene:
    level1_id: str
    level1_name: str
    level1_name_en: str
    scene: str
    elements: tuple[str, ...]


@dataclass(frozen=True)
class Sample:
    id: str
    image_path: str
    mask_path: str
    parent_group_id: str
    read_status: str = "ok"


@dataclass
class Review:
    quality_status: str = "unreviewed"
    scene_status: str = "unassigned"
    level1_id: str = ""
    level1_name: str = ""
    primary_level2_scene: str = ""
    secondary_scenes_json: str = "[]"
    scene_source: str = "taxonomy"
    custom_scene_name_en: str = ""
    custom_scene_name_zh: str = ""
    taxonomy_mapping_suggestion: str = ""
    issue_tags_json: str = "[]"
    reviewer_note: str = ""


@dataclass(frozen=True)
class ExportPlan:
    items: tuple[tuple[Sample, Review], ...]
    accepted_total: int
    assigned_total: int
    unassigned_total: int
    uncertain_total: int
    invalid_assigned_total: int


@dataclass(frozen=True)
class QualityImportEntry:
    source_row: int
    imported_sample_id: str
    image_name: str
    mask_name: str
    quality_status: str
    reviewer_note: str
    matched_sample_id: str = ""
    local_quality_status: str = ""
    local_reviewer_note: str = ""
    match_method: str = ""
    category: str = "invalid"
    detail: str = ""


@dataclass(frozen=True)
class QualityImportPlan:
    source_path: Path
    source_sha256: str
    sheet_name: str
    entries: tuple[QualityImportEntry, ...]

    def counts(self) -> Counter[str]:
        return Counter(entry.category for entry in self.entries)


def review_scene_name(review: Review) -> str:
    return review.primary_level2_scene or review.custom_scene_name_en


def build_export_plan(items: list[tuple[Sample, Review]], group_by_scene: bool) -> ExportPlan:
    accepted = [(sample, review) for sample, review in items if review.quality_status == "accepted"]
    assigned = [
        (sample, review)
        for sample, review in accepted
        if review.scene_status == "assigned" and bool(review_scene_name(review))
    ]
    return ExportPlan(
        items=tuple(assigned if group_by_scene else accepted),
        accepted_total=len(accepted),
        assigned_total=len(assigned),
        unassigned_total=sum(review.scene_status == "unassigned" for _, review in accepted),
        uncertain_total=sum(review.scene_status == "uncertain" for _, review in accepted),
        invalid_assigned_total=sum(
            review.scene_status == "assigned" and not review_scene_name(review) for _, review in accepted
        ),
    )


def export_review_items(plan: ExportPlan, root: Path, group_by_scene: bool,
                        image_root: Path | None = None, mask_root: Path | None = None) -> int:
    if plan.items:
        image_root = image_root or Path(os.path.commonpath([str(Path(s.image_path).resolve().parent) for s, _ in plan.items]))
        mask_root = mask_root or Path(os.path.commonpath([str(Path(s.mask_path).resolve().parent) for s, _ in plan.items]))
    destinations = set()
    for sample, review in plan.items:
        scene = review_scene_name(review)
        safe_scene = "".join(c if c.isalnum() or c in " ._-" else "_" for c in scene).strip() or "_unassigned"
        prefix = root / safe_scene if group_by_scene else root
        for source, source_root, kind in ((sample.image_path, image_root, "images"), (sample.mask_path, mask_root, "masks")):
            target = prefix / kind / Path(source).resolve().relative_to(source_root.resolve())
            key = str(target.resolve()).casefold()
            if key in destinations:
                raise ValueError(f"导出路径冲突：{target}")
            destinations.add(key)
    root.mkdir(parents=True, exist_ok=False)
    rows: list[list[str]] = []
    for sample, review in plan.items:
        scene = review_scene_name(review)
        export_root = root
        if group_by_scene:
            safe_scene = "".join(c if c.isalnum() or c in " ._-" else "_" for c in scene).strip() or "_unassigned"
            export_root = root / safe_scene
        image_dir = export_root / "images"
        mask_dir = export_root / "masks"
        image_dir.mkdir(parents=True, exist_ok=True)
        mask_dir.mkdir(parents=True, exist_ok=True)
        image_dst = image_dir / Path(sample.image_path).resolve().relative_to(image_root.resolve())
        mask_dst = mask_dir / Path(sample.mask_path).resolve().relative_to(mask_root.resolve())
        image_dst.parent.mkdir(parents=True, exist_ok=True)
        mask_dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(sample.image_path, image_dst)
        shutil.copy2(sample.mask_path, mask_dst)
        rows.append([
            sample.id,
            review.quality_status,
            review.scene_status,
            scene,
            review.scene_source,
            sample.image_path,
            sample.mask_path,
            str(image_dst),
            str(mask_dst),
            review.reviewer_note,
        ])
    with (root / "export_manifest.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([
            "sample_id",
            "quality_status",
            "scene_status",
            "scene",
            "scene_source",
            "source_image",
            "source_mask",
            "export_image",
            "export_mask",
            "reviewer_note",
        ])
        writer.writerows(rows)
    return len(rows)


QUALITY_NAMES = {"accepted": "合格", "rejected": "不合格", "needs_correction": "需修改"}
QUALITY_STATUS_BY_NAME = {
    "合格": "accepted",
    "不合格": "rejected",
    "需修改": "needs_correction",
    "accepted": "accepted",
    "rejected": "rejected",
    "needs_correction": "needs_correction",
}
QUALITY_IMPORT_CATEGORY_NAMES = {
    "ready": "可导入",
    "note_fill": "可补充备注",
    "same": "与本地相同",
    "conflict": "与本地冲突",
    "unmatched": "未匹配",
    "invalid": "无效记录",
    "duplicate": "重复记录",
    "duplicate_conflict": "表内冲突",
}


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _name_from_cell(value: Any) -> str:
    text = str(value or "").strip().replace("\\", "/")
    return text.rsplit("/", 1)[-1]


def _path_key(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix().casefold()
    except ValueError:
        return ""


def _suffix_path_candidates(
    value: Any, relative_index: dict[str, set[str]]
) -> tuple[set[str], str]:
    text = str(value or "").strip().replace("\\", "/").strip("/")
    parts = [part for part in text.split("/") if part and part != "."]
    for start in range(len(parts)):
        key = "/".join(parts[start:]).casefold()
        if key in relative_index:
            return set(relative_index[key]), key
    return set(), ""


def read_quality_progress_xlsx(
    path: Path,
    items: list[tuple[Sample, Review]],
    image_root: Path,
    mask_root: Path,
) -> QualityImportPlan:
    from openpyxl import load_workbook

    path = path.resolve()
    if path.suffix.casefold() != ".xlsx":
        raise ValueError("只支持读取 .xlsx 格式的质量审核记录。")
    if not path.is_file():
        raise FileNotFoundError(f"找不到 Excel 文件：{path}")
    if path.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Excel 文件超过 100 MB，请确认是否选择了正确的质量审核记录。")

    aliases = {
        "sample_id": {"样本id", "sampleid"},
        "image_name": {"原图名称", "imagename"},
        "mask_name": {"掩膜名称", "maskname"},
        "quality": {"质量类型", "qualitystatus"},
        "note": {"文本描述信息备注", "质量备注", "reviewernote"},
        "image_source": {"原图来源路径", "sourceimage"},
        "mask_source": {"掩膜来源路径", "sourcemask"},
        "image_export": {"原图导出路径", "exportimage"},
        "mask_export": {"掩膜导出路径", "exportmask"},
    }

    def normalized_header(value: Any) -> str:
        return "".join(str(value or "").strip().casefold().replace("_", "").split())

    samples_by_id = {sample.id: (sample, review) for sample, review in items}
    pair_index: dict[tuple[str, str], set[str]] = {}
    image_name_index: dict[str, set[str]] = {}
    image_relative_index: dict[str, set[str]] = {}
    mask_relative_index: dict[str, set[str]] = {}
    for sample, _review in items:
        image_name = Path(sample.image_path).name.casefold()
        mask_name = Path(sample.mask_path).name.casefold()
        pair_index.setdefault((image_name, mask_name), set()).add(sample.id)
        image_name_index.setdefault(image_name, set()).add(sample.id)
        image_relative = _path_key(Path(sample.image_path), image_root)
        mask_relative = _path_key(Path(sample.mask_path), mask_root)
        if image_relative:
            image_relative_index.setdefault(image_relative, set()).add(sample.id)
        if mask_relative:
            mask_relative_index.setdefault(mask_relative, set()).add(sample.id)

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["质量审核记录"] if "质量审核记录" in workbook.sheetnames else workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        header_lookup = {
            normalized_header(cell.value): index
            for index, cell in enumerate(header_cells)
            if normalized_header(cell.value)
        }
        columns: dict[str, int] = {}
        for field, field_aliases in aliases.items():
            for alias in field_aliases:
                if alias in header_lookup:
                    columns[field] = header_lookup[alias]
                    break
        if "quality" not in columns:
            raise ValueError("Excel 缺少“质量类型”列。")
        if not any(field in columns for field in ("sample_id", "image_name", "image_source", "image_export")):
            raise ValueError("Excel 缺少可用于识别样本的列。")

        entries: list[QualityImportEntry] = []
        supported_columns = tuple(columns.values())
        formula_fields = ("sample_id", "image_name", "mask_name", "quality", "image_source", "mask_source")
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if row_number > 200001:
                raise ValueError("Excel 有效范围超过 200000 行，已停止读取。")

            def value(field: str) -> Any:
                index = columns.get(field)
                return cells[index].value if index is not None and index < len(cells) else None

            if not any(
                index < len(cells) and cells[index].value not in (None, "")
                for index in supported_columns
            ):
                continue
            imported_id = str(value("sample_id") or "").strip()
            image_name = _name_from_cell(value("image_name"))
            mask_name = _name_from_cell(value("mask_name"))
            image_source = value("image_source") or value("image_export")
            mask_source = value("mask_source") or value("mask_export")
            image_name = image_name or _name_from_cell(image_source)
            mask_name = mask_name or _name_from_cell(mask_source)
            quality_text = str(value("quality") or "").strip()
            note = str(value("note") or "")

            formula_used = False
            for field in formula_fields:
                index = columns.get(field)
                if index is not None and index < len(cells) and cells[index].data_type == "f":
                    formula_used = True
                    break
            if formula_used:
                entries.append(QualityImportEntry(
                    row_number, imported_id, image_name, mask_name, "", note,
                    category="invalid", detail="识别字段或质量类型不能使用 Excel 公式",
                ))
                continue

            quality_status = QUALITY_STATUS_BY_NAME.get(quality_text.casefold(), "")
            if not quality_status:
                entries.append(QualityImportEntry(
                    row_number, imported_id, image_name, mask_name, "", note,
                    category="invalid", detail=f"无法识别质量类型：{quality_text or '空值'}",
                ))
                continue

            matched_ids: set[str] = set()
            match_method = ""
            if imported_id and imported_id in samples_by_id:
                sample = samples_by_id[imported_id][0]
                if image_name and image_name.casefold() != Path(sample.image_path).name.casefold():
                    entries.append(QualityImportEntry(
                        row_number, imported_id, image_name, mask_name, quality_status, note,
                        category="invalid", detail="样本 ID 存在，但原图名称与本地记录不一致",
                    ))
                    continue
                if mask_name and mask_name.casefold() != Path(sample.mask_path).name.casefold():
                    entries.append(QualityImportEntry(
                        row_number, imported_id, image_name, mask_name, quality_status, note,
                        category="invalid", detail="样本 ID 存在，但掩膜名称与本地记录不一致",
                    ))
                    continue
                matched_ids = {imported_id}
                match_method = "样本 ID"
            else:
                image_candidates, _ = _suffix_path_candidates(image_source, image_relative_index)
                mask_candidates, _ = _suffix_path_candidates(mask_source, mask_relative_index)
                if image_source and mask_source:
                    if image_candidates and mask_candidates:
                        matched_ids = image_candidates & mask_candidates
                        if matched_ids:
                            match_method = "相对路径"
                elif image_candidates:
                    matched_ids = image_candidates
                    match_method = "原图相对路径"
                elif mask_candidates:
                    matched_ids = mask_candidates
                    match_method = "掩膜相对路径"

                if not matched_ids and image_name and mask_name:
                    matched_ids = set(pair_index.get((image_name.casefold(), mask_name.casefold()), set()))
                    if matched_ids:
                        match_method = "原图名 + 掩膜名"
                elif not matched_ids and image_name and not mask_name:
                    matched_ids = set(image_name_index.get(image_name.casefold(), set()))
                    if matched_ids:
                        match_method = "唯一原图名"

            if len(matched_ids) != 1:
                detail = "本地工作区中没有找到对应样本"
                if len(matched_ids) > 1:
                    detail = f"匹配到 {len(matched_ids)} 个同名样本，无法自动确定"
                entries.append(QualityImportEntry(
                    row_number, imported_id, image_name, mask_name, quality_status, note,
                    category="unmatched", detail=detail,
                ))
                continue

            matched_id = next(iter(matched_ids))
            _sample, local_review = samples_by_id[matched_id]
            if local_review.quality_status == "unreviewed":
                category = "ready"
                detail = "本地未审核，可以导入"
            elif local_review.quality_status == quality_status and local_review.reviewer_note == note:
                category = "same"
                detail = "质量类型和备注均与本地相同"
            elif (
                local_review.quality_status == quality_status
                and not local_review.reviewer_note
                and bool(note)
            ):
                category = "note_fill"
                detail = "质量类型相同，可以补充本地空白备注"
            else:
                category = "conflict"
                detail = "本地质量类型或备注与导入记录不同"
            entries.append(QualityImportEntry(
                row_number, imported_id, image_name, mask_name, quality_status, note,
                matched_id, local_review.quality_status, local_review.reviewer_note,
                match_method, category, detail,
            ))
    finally:
        workbook.close()

    if not entries:
        raise ValueError("Excel 中没有可读取的质量审核记录。")

    indices_by_sample: dict[str, list[int]] = {}
    for index, entry in enumerate(entries):
        if entry.matched_sample_id and entry.category not in {"invalid", "unmatched"}:
            indices_by_sample.setdefault(entry.matched_sample_id, []).append(index)
    for matched_id, indices in indices_by_sample.items():
        if len(indices) < 2:
            continue
        values = {
            (entries[index].quality_status, entries[index].reviewer_note)
            for index in indices
        }
        if len(values) == 1:
            for index in indices[1:]:
                entries[index] = replace(
                    entries[index],
                    category="duplicate",
                    detail="Excel 中存在完全相同的重复记录，本行将跳过",
                )
        else:
            for index in indices:
                entries[index] = replace(
                    entries[index],
                    category="duplicate_conflict",
                    detail="Excel 中同一样本存在不同结果，必须先在表格中确认",
                )

    return QualityImportPlan(
        source_path=path,
        source_sha256=_file_sha256(path),
        sheet_name=sheet.title,
        entries=tuple(entries),
    )


def export_quality_items(
    items: list[tuple[Sample, Review]],
    root: Path,
    image_root: Path,
    mask_root: Path,
    statuses: set[str],
    copy_files: bool = True,
    include_table: bool = True,
    table_statuses: set[str] | None = None,
) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    file_statuses = statuses if copy_files else set()
    report_statuses = (statuses if table_statuses is None else table_statuses) if include_table else set()
    if not file_statuses and not report_statuses:
        raise ValueError("请至少选择一个导出项目。")
    selected = [
        (s, r) for s, r in items
        if r.quality_status in (file_statuses | report_statuses) and r.quality_status in QUALITY_NAMES
    ]
    if not selected:
        raise ValueError("所选导出项目没有已审核样本。")
    destinations: set[str] = set()
    report_samples: set[str] = set()
    planned = []
    for sample, review in selected:
        if review.quality_status in report_statuses:
            if sample.id in report_samples:
                raise ValueError(f"质量表格样本冲突：{sample.id}")
            report_samples.add(sample.id)
        sources = (Path(sample.image_path), Path(sample.mask_path))
        targets = []
        for source, source_root, kind in zip(sources, (image_root, mask_root), ("images", "masks")):
            relative = source.resolve().relative_to(source_root.resolve())
            target = root / QUALITY_NAMES[review.quality_status] / kind / relative
            key = str(target.resolve()).casefold()
            if review.quality_status in file_statuses:
                if key in destinations:
                    raise ValueError(f"导出路径冲突：{target}")
                destinations.add(key)
                if not source.is_file():
                    raise FileNotFoundError(f"找不到源文件：{source}")
            targets.append(target)
        planned.append((sample, review, sources, targets))
    root.mkdir(parents=True, exist_ok=False)
    workbook = Workbook() if include_table else None
    sheet = workbook.active if workbook else None
    if sheet is not None:
        sheet.title = "质量审核记录"
        sheet.append(["样本ID", "原图名称", "掩膜名称", "质量类型", "文本描述信息备注",
                      "场景名称", "原图来源路径", "掩膜来源路径", "原图导出路径", "掩膜导出路径"])
    try:
        for sample, review, sources, targets in planned:
            if review.quality_status in file_statuses:
                for source, target in zip(sources, targets):
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, target)
            if sheet is None or review.quality_status not in report_statuses:
                continue
            values = [sample.id, sources[0].name, sources[1].name,
                      QUALITY_NAMES[review.quality_status], review.reviewer_note,
                      review_scene_name(review), str(sources[0]), str(sources[1]),
                      str(targets[0].relative_to(root)) if review.quality_status in file_statuses else "",
                      str(targets[1].relative_to(root)) if review.quality_status in file_statuses else ""]
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet is not None:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column, width in zip("ABCDEFGHIJ", (24, 30, 36, 14, 60, 25, 50, 50, 50, 50)):
                sheet.column_dimensions[column].width = width
            workbook.save(root / "质量审核记录.xlsx")
    except Exception:
        shutil.rmtree(root)
        raise
    finally:
        if workbook is not None:
            workbook.close()
    return len(selected)

def now_text() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def normalized_path(path: Path) -> str:
    return str(path.resolve()).replace("\\", "/")


def resolve_export_start_directory(
    last_export_directory: Path | None,
    workspace_directory: Path | None,
    workspace_text: str = "",
    fallback: Path | None = None,
) -> str:
    candidates = [last_export_directory, workspace_directory]
    if workspace_text.strip():
        candidates.append(Path(workspace_text.strip()))
    for candidate in candidates:
        if candidate is not None and candidate.is_dir():
            return str(candidate)
    return str(fallback or Path.cwd())


def stable_sample_id(dataset_name: str, image_path: Path) -> str:
    stat = image_path.stat()
    raw = f"{dataset_name}|{normalized_path(image_path)}|{stat.st_size}".encode("utf-8")
    return hashlib.sha1(raw).hexdigest()[:16]


class Taxonomy:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path
        self.version = ""
        self.sha256 = ""
        self.scenes: list[Scene] = []
        self.by_domain: dict[str, list[Scene]] = {}
        if path is not None:
            self.load()

    def load(self) -> None:
        if self.path is None:
            return
        data = self.path.read_bytes()
        self.sha256 = hashlib.sha256(data).hexdigest()
        obj = json.loads(data.decode("utf-8"))
        self.version = str(obj.get("version", ""))
        scenes: list[Scene] = []
        for domain in obj.get("domains", []):
            for scene in domain.get("scenes", []):
                scenes.append(
                    Scene(
                        level1_id=str(domain.get("id", "")),
                        level1_name=str(domain.get("name_zh", "")),
                        level1_name_en=str(domain.get("name_en", "")),
                        scene=str(scene.get("scene", "")),
                        elements=tuple(scene.get("elements", [])),
                    )
                )
        self.scenes = scenes
        self.by_domain = {}
        for scene in scenes:
            self.by_domain.setdefault(scene.level1_id, []).append(scene)

    @property
    def domains(self) -> list[tuple[str, str, str]]:
        seen: dict[str, tuple[str, str, str]] = {}
        for scene in self.scenes:
            seen[scene.level1_id] = (scene.level1_id, scene.level1_name, scene.level1_name_en)
        return list(seen.values())

    def find_scene(self, name: str) -> Scene | None:
        for scene in self.scenes:
            if scene.scene == name:
                return scene
        return None


class ReviewDatabase:
    def __init__(self, db_path: Path) -> None:
        self.db_path = db_path
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(db_path))
        self.conn.row_factory = sqlite3.Row
        self.migrate()

    def migrate(self) -> None:
        cur = self.conn.cursor()
        cur.execute("PRAGMA journal_mode=WAL")
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS datasets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                image_root TEXT NOT NULL,
                mask_root TEXT NOT NULL,
                context_root TEXT,
                config_json TEXT NOT NULL,
                taxonomy_version TEXT,
                taxonomy_sha256 TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS samples (
                id TEXT PRIMARY KEY,
                dataset_id INTEGER NOT NULL,
                image_path TEXT NOT NULL,
                mask_path TEXT NOT NULL,
                context_path TEXT,
                parent_group_id TEXT,
                image_width INTEGER,
                image_height INTEGER,
                image_dtype TEXT,
                mask_width INTEGER,
                mask_height INTEGER,
                read_status TEXT,
                content_sha256 TEXT
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS reviews (
                sample_id TEXT PRIMARY KEY,
                quality_status TEXT NOT NULL,
                scene_status TEXT NOT NULL,
                level1_id TEXT,
                level1_name TEXT,
                primary_level2_scene TEXT,
                secondary_scenes_json TEXT,
                scene_source TEXT,
                custom_scene_name_en TEXT,
                custom_scene_name_zh TEXT,
                taxonomy_mapping_suggestion TEXT,
                issue_tags_json TEXT,
                reviewer_note TEXT,
                reviewed_at TEXT,
                updated_at TEXT
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS dataset_scene_vocab (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_id INTEGER NOT NULL,
                scene_source TEXT NOT NULL,
                scene_name_en TEXT NOT NULL,
                scene_name_zh TEXT,
                level1_id TEXT,
                level1_name TEXT,
                mapped_level2_scene TEXT,
                is_pinned INTEGER DEFAULT 0,
                use_count INTEGER DEFAULT 0,
                last_used_at TEXT,
                created_at TEXT NOT NULL,
                UNIQUE(dataset_id, scene_source, scene_name_en)
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS quality_imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_id INTEGER NOT NULL,
                source_path TEXT NOT NULL,
                source_sha256 TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                policy TEXT NOT NULL,
                total_rows INTEGER NOT NULL,
                applied_count INTEGER NOT NULL,
                conflict_count INTEGER NOT NULL,
                unmatched_count INTEGER NOT NULL,
                invalid_count INTEGER NOT NULL,
                imported_at TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS quality_import_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                import_id INTEGER NOT NULL,
                source_row INTEGER NOT NULL,
                sample_id TEXT,
                imported_quality_status TEXT,
                imported_reviewer_note TEXT,
                local_quality_status TEXT,
                local_reviewer_note TEXT,
                category TEXT NOT NULL,
                outcome TEXT NOT NULL,
                detail TEXT
            )
            """
        )
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()

    def create_dataset(
        self,
        name: str,
        image_root: Path,
        mask_root: Path,
        config: dict[str, Any],
        taxonomy: Taxonomy,
    ) -> int:
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO datasets
            (name, image_root, mask_root, context_root, config_json, taxonomy_version, taxonomy_sha256, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name,
                normalized_path(image_root),
                normalized_path(mask_root),
                "",
                json.dumps(config, ensure_ascii=False),
                taxonomy.version,
                taxonomy.sha256,
                now_text(),
            ),
        )
        self.conn.commit()
        return int(cur.lastrowid)

    def latest_dataset_id(self) -> int | None:
        row = self.conn.execute("SELECT id FROM datasets ORDER BY id DESC LIMIT 1").fetchone()
        return int(row["id"]) if row else None

    def latest_dataset(self) -> sqlite3.Row | None:
        return self.conn.execute("SELECT * FROM datasets ORDER BY id DESC LIMIT 1").fetchone()

    def add_samples(self, samples: list[Sample], dataset_id: int) -> None:
        rows = [
            (
                s.id,
                dataset_id,
                s.image_path,
                s.mask_path,
                "",
                s.parent_group_id,
                None,
                None,
                "",
                None,
                None,
                s.read_status,
                "",
            )
            for s in samples
        ]
        self.conn.executemany(
            """
            INSERT OR IGNORE INTO samples
            (id, dataset_id, image_path, mask_path, context_path, parent_group_id, image_width, image_height,
             image_dtype, mask_width, mask_height, read_status, content_sha256)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            rows,
        )
        self.conn.commit()

    def samples(self, dataset_id: int, status_filter: str = "all", text_filter: str = "") -> list[tuple[Sample, Review]]:
        sql = """
            SELECT s.*, r.quality_status, r.scene_status, r.level1_id, r.level1_name,
                   r.primary_level2_scene, r.secondary_scenes_json, r.scene_source,
                   r.custom_scene_name_en, r.custom_scene_name_zh, r.taxonomy_mapping_suggestion,
                   r.issue_tags_json, r.reviewer_note
            FROM samples s
            LEFT JOIN reviews r ON r.sample_id = s.id
            WHERE s.dataset_id = ?
        """
        params: list[Any] = [dataset_id]
        if status_filter == "unreviewed":
            sql += " AND COALESCE(r.quality_status, 'unreviewed') = 'unreviewed'"
        elif status_filter == "accepted":
            sql += " AND COALESCE(r.quality_status, 'unreviewed') = 'accepted'"
        elif status_filter == "needs_correction":
            sql += " AND COALESCE(r.quality_status, 'unreviewed') = 'needs_correction'"
        elif status_filter == "rejected":
            sql += " AND COALESCE(r.quality_status, 'unreviewed') = 'rejected'"
        elif status_filter == "unassigned":
            sql += " AND COALESCE(r.scene_status, 'unassigned') = 'unassigned'"
        if text_filter.strip():
            sql += " AND (s.image_path LIKE ? OR s.id LIKE ? OR COALESCE(r.reviewer_note, '') LIKE ?)"
            like = f"%{text_filter.strip()}%"
            params.extend([like, like, like])
        sql += " ORDER BY s.parent_group_id, s.image_path"
        result: list[tuple[Sample, Review]] = []
        for row in self.conn.execute(sql, params).fetchall():
            sample = Sample(
                id=row["id"],
                image_path=row["image_path"],
                mask_path=row["mask_path"],
                parent_group_id=row["parent_group_id"] or "",
                read_status=row["read_status"] or "ok",
            )
            review = Review(
                quality_status=row["quality_status"] or "unreviewed",
                scene_status=row["scene_status"] or "unassigned",
                level1_id=row["level1_id"] or "",
                level1_name=row["level1_name"] or "",
                primary_level2_scene=row["primary_level2_scene"] or "",
                secondary_scenes_json=row["secondary_scenes_json"] or "[]",
                scene_source=row["scene_source"] or "taxonomy",
                custom_scene_name_en=row["custom_scene_name_en"] or "",
                custom_scene_name_zh=row["custom_scene_name_zh"] or "",
                taxonomy_mapping_suggestion=row["taxonomy_mapping_suggestion"] or "",
                issue_tags_json=row["issue_tags_json"] or "[]",
                reviewer_note=row["reviewer_note"] or "",
            )
            result.append((sample, review))
        return result

    def save_review(self, sample_id: str, dataset_id: int, review: Review) -> None:
        review.updated_at = now_text() if hasattr(review, "updated_at") else now_text()
        reviewed_at = now_text() if review.quality_status != "unreviewed" or review.scene_status != "unassigned" else ""
        self.conn.execute(
            """
            INSERT INTO reviews
            (sample_id, quality_status, scene_status, level1_id, level1_name, primary_level2_scene,
             secondary_scenes_json, scene_source, custom_scene_name_en, custom_scene_name_zh,
             taxonomy_mapping_suggestion, issue_tags_json, reviewer_note, reviewed_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(sample_id) DO UPDATE SET
              quality_status=excluded.quality_status,
              scene_status=excluded.scene_status,
              level1_id=excluded.level1_id,
              level1_name=excluded.level1_name,
              primary_level2_scene=excluded.primary_level2_scene,
              secondary_scenes_json=excluded.secondary_scenes_json,
              scene_source=excluded.scene_source,
              custom_scene_name_en=excluded.custom_scene_name_en,
              custom_scene_name_zh=excluded.custom_scene_name_zh,
              taxonomy_mapping_suggestion=excluded.taxonomy_mapping_suggestion,
              issue_tags_json=excluded.issue_tags_json,
              reviewer_note=excluded.reviewer_note,
              reviewed_at=excluded.reviewed_at,
              updated_at=excluded.updated_at
            """,
            (
                sample_id,
                review.quality_status,
                review.scene_status,
                review.level1_id,
                review.level1_name,
                review.primary_level2_scene,
                review.secondary_scenes_json,
                review.scene_source,
                review.custom_scene_name_en,
                review.custom_scene_name_zh,
                review.taxonomy_mapping_suggestion,
                review.issue_tags_json,
                review.reviewer_note,
                reviewed_at,
                now_text(),
            ),
        )
        self._touch_scene_vocab(dataset_id, review)
        self.conn.commit()

    def quality_import_seen(self, dataset_id: int, source_sha256: str) -> bool:
        row = self.conn.execute(
            "SELECT 1 FROM quality_imports WHERE dataset_id = ? AND source_sha256 = ? LIMIT 1",
            (dataset_id, source_sha256),
        ).fetchone()
        return row is not None

    def backup_to(self, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        backup = sqlite3.connect(str(destination))
        try:
            self.conn.backup(backup)
        finally:
            backup.close()

    def apply_quality_import(
        self,
        dataset_id: int,
        plan: QualityImportPlan,
        policy: str,
        conflict_overrides: dict[str, str] | None = None,
    ) -> tuple[int, int]:
        allowed = {
            "fill_unreviewed": {"ready"},
            "keep_local": {"ready", "note_fill"},
            "use_imported": {"ready", "note_fill", "conflict"},
        }
        if policy not in allowed:
            raise ValueError(f"未知的导入策略：{policy}")
        applicable = allowed[policy]
        conflict_overrides = conflict_overrides or {}
        if any(action not in {"keep_local", "use_imported"} for action in conflict_overrides.values()):
            raise ValueError("存在未知的逐条冲突处理方式。")
        imported_at = now_text()
        def should_apply(entry: QualityImportEntry) -> bool:
            if entry.category == "conflict":
                action = conflict_overrides.get(entry.matched_sample_id)
                if action:
                    return action == "use_imported"
            return entry.category in applicable

        applied_count = sum(should_apply(entry) for entry in plan.entries)
        counts = plan.counts()
        try:
            self.conn.execute("BEGIN")
            cursor = self.conn.execute(
                """
                INSERT INTO quality_imports
                (dataset_id, source_path, source_sha256, sheet_name, policy, total_rows,
                 applied_count, conflict_count, unmatched_count, invalid_count, imported_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    dataset_id,
                    str(plan.source_path),
                    plan.source_sha256,
                    plan.sheet_name,
                    policy,
                    len(plan.entries),
                    applied_count,
                    counts["conflict"] + counts["duplicate_conflict"],
                    counts["unmatched"],
                    counts["invalid"],
                    imported_at,
                ),
            )
            import_id = int(cursor.lastrowid)
            for entry in plan.entries:
                apply_entry = should_apply(entry)
                outcome = "applied" if apply_entry else "skipped"
                if apply_entry:
                    self.conn.execute(
                        """
                        INSERT INTO reviews
                        (sample_id, quality_status, scene_status, reviewer_note, reviewed_at, updated_at)
                        VALUES (?, ?, 'unassigned', ?, ?, ?)
                        ON CONFLICT(sample_id) DO UPDATE SET
                          quality_status=excluded.quality_status,
                          reviewer_note=excluded.reviewer_note,
                          reviewed_at=excluded.reviewed_at,
                          updated_at=excluded.updated_at
                        """,
                        (
                            entry.matched_sample_id,
                            entry.quality_status,
                            entry.reviewer_note,
                            imported_at,
                            imported_at,
                        ),
                    )
                self.conn.execute(
                    """
                    INSERT INTO quality_import_items
                    (import_id, source_row, sample_id, imported_quality_status,
                     imported_reviewer_note, local_quality_status, local_reviewer_note,
                     category, outcome, detail)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        import_id,
                        entry.source_row,
                        entry.matched_sample_id,
                        entry.quality_status,
                        entry.reviewer_note,
                        entry.local_quality_status,
                        entry.local_reviewer_note,
                        entry.category,
                        outcome,
                        entry.detail,
                    ),
                )
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return import_id, applied_count

    def _touch_scene_vocab(self, dataset_id: int, review: Review) -> None:
        scene_name = review.primary_level2_scene
        scene_zh = review.primary_level2_scene
        mapped = review.primary_level2_scene
        if review.scene_source != "taxonomy":
            scene_name = review.custom_scene_name_en.strip()
            scene_zh = review.custom_scene_name_zh.strip()
            mapped = review.taxonomy_mapping_suggestion.strip()
        if not scene_name:
            return
        self.conn.execute(
            """
            INSERT INTO dataset_scene_vocab
            (dataset_id, scene_source, scene_name_en, scene_name_zh, level1_id, level1_name,
             mapped_level2_scene, is_pinned, use_count, last_used_at, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, 0, 1, ?, ?)
            ON CONFLICT(dataset_id, scene_source, scene_name_en) DO UPDATE SET
              scene_name_zh=excluded.scene_name_zh,
              level1_id=excluded.level1_id,
              level1_name=excluded.level1_name,
              mapped_level2_scene=excluded.mapped_level2_scene,
              last_used_at=excluded.last_used_at
            """,
            (
                dataset_id,
                review.scene_source,
                scene_name,
                scene_zh,
                review.level1_id,
                review.level1_name,
                mapped,
                now_text(),
                now_text(),
            ),
        )

    def common_scenes(self, dataset_id: int) -> list[sqlite3.Row]:
        return self.conn.execute(
            """
            WITH assigned AS (
                SELECT
                    COALESCE(scene_source, 'taxonomy') AS scene_source,
                    CASE
                        WHEN COALESCE(scene_source, 'taxonomy') = 'taxonomy'
                        THEN primary_level2_scene
                        ELSE custom_scene_name_en
                    END AS scene_name_en,
                    COUNT(*) AS assigned_count
                FROM reviews
                WHERE sample_id IN (SELECT id FROM samples WHERE dataset_id = ?)
                  AND scene_status IN ('assigned', 'uncertain')
                  AND COALESCE(
                        CASE
                            WHEN COALESCE(scene_source, 'taxonomy') = 'taxonomy'
                            THEN primary_level2_scene
                            ELSE custom_scene_name_en
                        END,
                        ''
                      ) <> ''
                GROUP BY scene_source, scene_name_en
            )
            SELECT v.*, COALESCE(a.assigned_count, 0) AS assigned_count
            FROM dataset_scene_vocab v
            LEFT JOIN assigned a
              ON a.scene_source = v.scene_source
             AND a.scene_name_en = v.scene_name_en
            WHERE v.dataset_id = ?
            ORDER BY v.is_pinned DESC, assigned_count DESC, v.last_used_at DESC, v.scene_name_en
            """,
            (dataset_id, dataset_id),
        ).fetchall()

    def stats(self, dataset_id: int) -> dict[str, Any]:
        total = self.conn.execute("SELECT COUNT(*) AS n FROM samples WHERE dataset_id = ?", (dataset_id,)).fetchone()["n"]
        quality = self.conn.execute(
            """
            SELECT COALESCE(r.quality_status, 'unreviewed') AS status, COUNT(*) AS n
            FROM samples s LEFT JOIN reviews r ON r.sample_id = s.id
            WHERE s.dataset_id = ?
            GROUP BY COALESCE(r.quality_status, 'unreviewed')
            """,
            (dataset_id,),
        ).fetchall()
        scenes = self.conn.execute(
            """
            SELECT COALESCE(NULLIF(r.primary_level2_scene, ''), NULLIF(r.custom_scene_name_en, ''), '_unassigned') AS scene,
                   COALESCE(r.scene_source, 'taxonomy') AS source,
                   COUNT(*) AS n
            FROM samples s LEFT JOIN reviews r ON r.sample_id = s.id
            WHERE s.dataset_id = ?
            GROUP BY scene, source
            ORDER BY n DESC, scene
            """,
            (dataset_id,),
        ).fetchall()
        return {
            "total": total,
            "quality": {row["status"]: row["n"] for row in quality},
            "scenes": [(row["scene"], row["source"], row["n"]) for row in scenes],
        }


def scan_dataset(
    dataset_name: str,
    image_root: Path,
    mask_root: Path,
    mode: str,
    image_dir_name: str,
    mask_dir_name: str,
    split_name: str = "all",
    recursive: bool = True,
    mask_name_prefix: str = "",
    mask_name_suffix: str = "",
) -> tuple[list[Sample], list[str]]:
    samples: list[Sample] = []
    warnings: list[str] = []
    mask_name_prefix = mask_name_prefix.strip()
    mask_name_suffix = mask_name_suffix.strip()

    def normalized_mask_stem(stem: str) -> str:
        if mask_name_prefix and stem.startswith(mask_name_prefix):
            stem = stem[len(mask_name_prefix) :]
        if mask_name_suffix and stem.endswith(mask_name_suffix):
            stem = stem[: -len(mask_name_suffix)]
        return stem

    def image_pairing_key(path: Path, root: Path) -> str:
        if mode == "relative_path_stem":
            return path.relative_to(root).with_suffix("").as_posix().lower()
        return path.stem.lower()

    def mask_pairing_key(path: Path, root: Path) -> str:
        stem = normalized_mask_stem(path.stem)
        if mode == "relative_path_stem":
            return path.relative_to(root).with_name(stem).with_suffix("").as_posix().lower()
        return stem.lower()

    def image_named_mask_candidate(image_path: Path) -> Path:
        mask_stem = f"{mask_name_prefix}{image_path.stem}{mask_name_suffix}"
        return image_path.with_name(f"{mask_stem}{image_path.suffix}")

    if mode == "voc_segmentation":
        image_dir = image_root / image_dir_name
        mask_dir = mask_root / mask_dir_name
        if not image_dir.is_dir():
            image_dir = image_root
        if not mask_dir.is_dir():
            mask_dir = mask_root

        image_by_stem = {
            path.stem: path
            for path in image_dir.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_EXTS
        }
        mask_by_stem = {
            normalized_mask_stem(path.stem): path
            for path in mask_dir.iterdir()
            if path.is_file() and path.suffix.lower() in MASK_EXTS
        }

        split = split_name.strip().lower() or "all"
        if split == "all":
            sample_ids = sorted(image_by_stem)
        else:
            split_file = image_root / "ImageSets" / "Segmentation" / f"{split}.txt"
            if not split_file.is_file():
                raise ValueError(f"VOC 划分文件不存在: {split_file}")
            sample_ids = [
                line.strip().split()[0]
                for line in split_file.read_text(encoding="utf-8-sig").splitlines()
                if line.strip()
            ]

        seen_ids: set[str] = set()
        for sample_name in sample_ids:
            if sample_name in seen_ids:
                warnings.append(f"duplicate sample in VOC split: {sample_name}")
                continue
            seen_ids.add(sample_name)
            image_path = image_by_stem.get(sample_name)
            mask_path = mask_by_stem.get(sample_name)
            if image_path is None:
                warnings.append(f"missing image for VOC id: {sample_name}")
                continue
            if mask_path is None:
                warnings.append(f"missing mask for VOC id: {sample_name}")
                continue
            samples.append(
                Sample(
                    id=stable_sample_id(dataset_name, image_path),
                    image_path=normalized_path(image_path),
                    mask_path=normalized_path(mask_path),
                    parent_group_id=split,
                )
            )
    elif mode == "mirrored_subdirectories":
        image_files = [p for p in image_root.rglob("*") if p.is_file() and p.suffix.lower() in IMAGE_EXTS and image_dir_name in p.parts]
        for image_path in image_files:
            parts = list(image_path.parts)
            try:
                idx = len(parts) - 1 - parts[::-1].index(image_dir_name)
            except ValueError:
                continue
            parts[idx] = mask_dir_name
            candidate = Path(*parts)
            if mask_name_prefix or mask_name_suffix:
                candidate = image_named_mask_candidate(candidate)
            mask_path = find_mask_with_extensions(candidate, MASK_EXTS)
            if mask_path is None:
                warnings.append(f"missing mask: {image_path}")
                continue
            parent = image_path.parent.parent.name if image_path.parent.parent else ""
            samples.append(
                Sample(
                    id=stable_sample_id(dataset_name, image_path),
                    image_path=normalized_path(image_path),
                    mask_path=normalized_path(mask_path),
                    parent_group_id=parent,
                )
            )
    elif mode == "deepglobe_suffix":
        image_files = deepglobe_role_files(
            image_root, IMAGE_EXTS, DEEPGLOBE_IMAGE_SUFFIX, recursive
        )
        mask_files = deepglobe_role_files(
            mask_root, MASK_EXTS, DEEPGLOBE_MASK_SUFFIX, recursive
        )

        def deepglobe_key(path: Path, root: Path, suffix: str) -> str:
            relative_path = path.relative_to(root)
            stem = relative_path.stem
            if not stem.lower().endswith(suffix):
                raise ValueError(f"DeepGlobe 文件名缺少后缀 {suffix}: {path}")
            normalized_stem = stem[: -len(suffix)]
            return (relative_path.parent / normalized_stem).as_posix().lower()

        images_by_key: dict[str, list[Path]] = {}
        masks_by_key: dict[str, list[Path]] = {}
        for image_path in image_files:
            key = deepglobe_key(image_path, image_root, DEEPGLOBE_IMAGE_SUFFIX)
            images_by_key.setdefault(key, []).append(image_path)
        for mask_path in mask_files:
            key = deepglobe_key(mask_path, mask_root, DEEPGLOBE_MASK_SUFFIX)
            masks_by_key.setdefault(key, []).append(mask_path)

        duplicate_keys = {
            key for key, paths in images_by_key.items() if len(paths) > 1
        } | {
            key for key, paths in masks_by_key.items() if len(paths) > 1
        }
        for key in sorted(duplicate_keys):
            warnings.append(f"ambiguous DeepGlobe pairing key: {key}")

        for key, paths in images_by_key.items():
            if key in duplicate_keys:
                continue
            mask_candidates = masks_by_key.get(key, [])
            if not mask_candidates:
                warnings.append(f"missing DeepGlobe mask for id: {key}")
                continue
            image_path = paths[0]
            samples.append(
                Sample(
                    id=stable_sample_id(dataset_name, image_path),
                    image_path=normalized_path(image_path),
                    mask_path=normalized_path(mask_candidates[0]),
                    parent_group_id=image_path.parent.name,
                )
            )
        for key, paths in masks_by_key.items():
            if key not in images_by_key and key not in duplicate_keys:
                for _mask_path in paths:
                    warnings.append(f"missing DeepGlobe image for id: {key}")
    else:
        image_iter = image_root.rglob("*") if recursive else image_root.iterdir()
        mask_iter = mask_root.rglob("*") if recursive else mask_root.iterdir()
        image_files = sorted(
            path for path in image_iter if path.is_file() and path.suffix.lower() in IMAGE_EXTS
        )
        mask_files = sorted(
            path for path in mask_iter if path.is_file() and path.suffix.lower() in MASK_EXTS
        )

        images_by_key: dict[str, list[Path]] = {}
        masks_by_key: dict[str, list[Path]] = {}
        for image_path in image_files:
            images_by_key.setdefault(image_pairing_key(image_path, image_root), []).append(image_path)
        for mask_path in mask_files:
            masks_by_key.setdefault(mask_pairing_key(mask_path, mask_root), []).append(mask_path)

        duplicate_keys = {
            key for key, paths in images_by_key.items() if len(paths) > 1
        } | {
            key for key, paths in masks_by_key.items() if len(paths) > 1
        }
        for key in sorted(duplicate_keys):
            warnings.append(f"ambiguous duplicate pairing key: {key}")

        for image_path in image_files:
            key = image_pairing_key(image_path, image_root)
            if key in duplicate_keys:
                continue
            mask_candidates = masks_by_key.get(key, [])
            mask_path = mask_candidates[0] if mask_candidates else None
            if mask_path is None:
                warnings.append(f"missing mask: {image_path}")
                continue
            parent = image_path.parent.name
            samples.append(
                Sample(
                    id=stable_sample_id(dataset_name, image_path),
                    image_path=normalized_path(image_path),
                    mask_path=normalized_path(mask_path),
                    parent_group_id=parent,
                )
            )
        for key, paths in masks_by_key.items():
            if key not in images_by_key:
                for mask_path in paths:
                    warnings.append(f"orphan mask: {mask_path}")
    return samples, warnings


def supported_files(root: Path, extensions: set[str], recursive: bool = True) -> list[Path]:
    iterator = root.rglob("*") if recursive else root.iterdir()
    return sorted(path for path in iterator if path.is_file() and path.suffix.lower() in extensions)


def deepglobe_role_files(
    root: Path,
    extensions: set[str],
    stem_suffix: str,
    recursive: bool = True,
) -> list[Path]:
    """Return files for one DeepGlobe role when images and masks share a folder."""
    return [
        path
        for path in supported_files(root, extensions, recursive)
        if path.stem.lower().endswith(stem_suffix)
    ]


LABEL_COLORS = (
    "#00d7ff",
    "#ffb000",
    "#7dde92",
    "#ff6b8a",
    "#a78bfa",
    "#f97316",
    "#22c55e",
    "#38bdf8",
)
INDEXED_MASK_MODES = {"1", "L", "P", "I", "I;16", "I;16B", "I;16L"}


def mask_encoding(mask: Image.Image) -> str:
    return "indexed" if mask.mode in INDEXED_MASK_MODES else "rgb"


def mask_label_key(encoding: str, value: int | tuple[int, int, int]) -> str:
    if encoding == "indexed":
        return f"i:{int(value)}"
    red, green, blue = value
    return f"rgb:{red},{green},{blue}"


def mask_label_text(value: int | list[int] | tuple[int, int, int]) -> str:
    if isinstance(value, int):
        return str(value)
    return f"({', '.join(str(int(channel)) for channel in value)})"


def default_label_color(
    encoding: str,
    value: int | tuple[int, int, int],
    palette: list[int] | None = None,
) -> str:
    if encoding == "rgb":
        red, green, blue = value
        return f"#{red:02x}{green:02x}{blue:02x}"
    index = int(value)
    if palette and 0 <= index * 3 + 2 < len(palette):
        red, green, blue = palette[index * 3:index * 3 + 3]
        return f"#{red:02x}{green:02x}{blue:02x}"
    if index == 0:
        return "#000000"
    digest = int(hashlib.sha1(str(index).encode("ascii")).hexdigest()[:8], 16)
    return LABEL_COLORS[digest % len(LABEL_COLORS)]


def _edge_label_counts(mask: Image.Image, encoding: str) -> Counter[str]:
    source = mask if encoding == "indexed" else mask.convert("RGB")
    array = np.asarray(source)
    if array.ndim == 2:
        edge = np.concatenate((array[0, :], array[-1, :], array[:, 0], array[:, -1]))
        values, counts = np.unique(edge, return_counts=True)
        return Counter({mask_label_key(encoding, int(value)): int(count) for value, count in zip(values, counts)})
    edge = np.concatenate((array[0, :, :3], array[-1, :, :3], array[:, 0, :3], array[:, -1, :3]))
    values, counts = np.unique(edge.reshape(-1, 3), axis=0, return_counts=True)
    return Counter(
        {
            mask_label_key(encoding, tuple(int(channel) for channel in value)): int(count)
            for value, count in zip(values, counts)
        }
    )


def infer_mask_schema(inspection: dict[str, Any], default_foreground_name: str = "") -> dict[str, Any]:
    detected_labels = inspection.get("mask_labels", [])
    if not detected_labels:
        return {"schema_version": 3, "encoding": inspection.get("mask_encoding", "unknown"), "background_mode": "none", "labels": []}
    by_key = {label["key"]: label for label in detected_labels}
    preferred_backgrounds = ("i:0", "rgb:0,0,0")
    background_key = next((key for key in preferred_backgrounds if key in by_key), "")
    if not background_key:
        background_key = max(
            detected_labels,
            key=lambda label: (label.get("border_count", 0), label.get("pixel_count", 0)),
        )["key"]
    class_count = len(detected_labels) - 1
    labels: list[dict[str, Any]] = []
    for label in detected_labels:
        role = "background" if label["key"] == background_key else "class"
        if role == "background":
            name = "background"
        elif class_count == 1 and default_foreground_name.strip():
            name = default_foreground_name.strip()
        elif class_count == 1:
            name = "foreground"
        else:
            safe_value = mask_label_text(label["value"]).replace("(", "").replace(")", "").replace(", ", "_")
            name = f"class_{safe_value}"
        labels.append(
            {
                "key": label["key"],
                "value": label["value"],
                "role": role,
                "name": name,
                "color": label["color"],
            }
        )
    return {
        "schema_version": 3,
        "encoding": inspection.get("mask_encoding", "unknown"),
        "background_mode": "explicit",
        "labels": labels,
    }


def validate_mask_schema(schema: dict[str, Any]) -> None:
    labels = schema.get("labels", [])
    if not labels:
        raise ValueError("Mask 类别映射为空。")
    if any(label.get("role") not in {"background", "class", "ignore"} for label in labels):
        raise ValueError("Mask 类别映射中存在无效角色。")
    classes = [label for label in labels if label.get("role") == "class"]
    if not classes:
        raise ValueError("请在 Mask 类别映射中至少指定一个有效类别。")
    for label in classes:
        if not str(label.get("name") or "").strip():
            raise ValueError(f"请为标签 {mask_label_text(label['value'])} 填写类别名称。")


def inspect_sample_pairs(samples: list[Sample], limit: int = 64) -> dict[str, Any]:
    if not samples:
        return {
            "checked": 0,
            "mask_values": [],
            "mask_labels": [],
            "mask_encoding": "unknown",
            "mask_modes": [],
            "label_overflow": 0,
            "size_mismatches": 0,
            "read_errors": 0,
        }
    count = min(limit, len(samples))
    if count == 1:
        selected = [samples[0]]
    else:
        indices = {round(index * (len(samples) - 1) / (count - 1)) for index in range(count)}
        selected = [samples[index] for index in sorted(indices)]

    labels: dict[str, dict[str, Any]] = {}
    encodings: set[str] = set()
    modes: set[str] = set()
    border_counts: Counter[str] = Counter()
    label_overflow = 0
    size_mismatches = 0
    read_errors = 0
    for sample in selected:
        try:
            with Image.open(sample.image_path) as image, Image.open(sample.mask_path) as mask:
                if image.size != mask.size:
                    size_mismatches += 1
                encoding = mask_encoding(mask)
                encodings.add(encoding)
                modes.add(mask.mode)
                palette = mask.getpalette() if mask.mode == "P" else None
                source = mask if encoding == "indexed" else mask.convert("RGB")
                colors = source.getcolors(maxcolors=4096)
                if colors is not None:
                    seen_in_sample: set[str] = set()
                    for pixel_count, raw_value in colors:
                        if encoding == "indexed":
                            value: int | tuple[int, int, int] = int(raw_value)
                        else:
                            value = tuple(int(channel) for channel in raw_value[:3])
                        key = mask_label_key(encoding, value)
                        if key not in labels:
                            labels[key] = {
                                "key": key,
                                "value": value,
                                "color": default_label_color(encoding, value, palette),
                                "pixel_count": 0,
                                "sample_count": 0,
                            }
                        labels[key]["pixel_count"] += int(pixel_count)
                        seen_in_sample.add(key)
                    for key in seen_in_sample:
                        labels[key]["sample_count"] += 1
                else:
                    label_overflow += 1
                border_counts.update(_edge_label_counts(mask, encoding))
        except Exception:
            read_errors += 1
    for key, label in labels.items():
        label["border_count"] = border_counts[key]
    ordered_labels = sorted(labels.values(), key=lambda label: (-label["pixel_count"], label["key"]))
    indexed_values = sorted(
        int(label["value"]) for label in ordered_labels if label["key"].startswith("i:")
    )
    return {
        "checked": len(selected),
        "mask_values": indexed_values,
        "mask_labels": ordered_labels,
        "mask_encoding": next(iter(encodings)) if len(encodings) == 1 else "mixed",
        "mask_modes": sorted(modes),
        "label_overflow": label_overflow,
        "size_mismatches": size_mismatches,
        "read_errors": read_errors,
    }


def find_mask_with_extensions(path: Path, exts: set[str]) -> Path | None:
    if path.exists():
        return path
    for ext in exts:
        candidate = path.with_suffix(ext)
        if candidate.exists():
            return candidate
    return None


def pil_to_qimage(image: Image.Image) -> QImage:
    rgba = image.convert("RGBA")
    data = rgba.tobytes("raw", "RGBA")
    qimage = QImage(data, rgba.width, rgba.height, rgba.width * 4, QImage.Format.Format_RGBA8888)
    return qimage.copy()


def pil_to_qpixmap(image: Image.Image) -> QPixmap:
    return QPixmap.fromImage(pil_to_qimage(image))


def _label_code(encoding: str, value: int | list[int] | tuple[int, int, int]) -> int:
    if encoding == "indexed":
        return int(value)
    red, green, blue = (int(channel) for channel in value[:3])
    return (red << 16) | (green << 8) | blue


def compose_preview_layers(
    image_path: str,
    mask_path: str,
    foreground_values: set[int] | None = None,
    mask_schema: dict[str, Any] | None = None,
    visible_label_keys: set[str] | None = None,
) -> tuple[Image.Image, Image.Image]:
    with Image.open(image_path) as source_image:
        image = source_image.convert("RGB")
    with Image.open(mask_path) as mask:
        if mask.size != image.size:
            raise ValueError(f"image and mask size mismatch: {image.size} vs {mask.size}")
        overlay_array = np.zeros((image.height, image.width, 4), dtype=np.uint8)
        if mask_schema and mask_schema.get("labels"):
            encoding = mask_encoding(mask)
            source = mask if encoding == "indexed" else mask.convert("RGB")
            source_array = np.asarray(source)
            if encoding == "indexed":
                pixel_codes = source_array.astype(np.uint32, copy=False)
            else:
                rgb = source_array[:, :, :3].astype(np.uint32, copy=False)
                pixel_codes = (rgb[:, :, 0] << 16) | (rgb[:, :, 1] << 8) | rgb[:, :, 2]
            active_labels = [
                label
                for label in mask_schema["labels"]
                if label.get("role") == "class"
                and (visible_label_keys is None or label.get("key") in visible_label_keys)
            ]
            entries: list[tuple[int, tuple[int, int, int, int]]] = []
            for label in active_labels:
                color = str(label.get("color") or "#00d7ff").lstrip("#")
                try:
                    channels = tuple(int(color[index:index + 2], 16) for index in (0, 2, 4))
                    if len(channels) != 3:
                        raise ValueError
                except (ValueError, TypeError):
                    channels = (0, 215, 255)
                entries.append((_label_code(encoding, label["value"]), (*channels, 255)))
            if entries:
                entries.sort(key=lambda entry: entry[0])
                lookup_codes = np.asarray([entry[0] for entry in entries], dtype=np.uint32)
                lookup_colors = np.asarray([entry[1] for entry in entries], dtype=np.uint8)
                flat_codes = pixel_codes.reshape(-1)
                positions = np.searchsorted(lookup_codes, flat_codes)
                safe_positions = np.minimum(positions, len(lookup_codes) - 1)
                matched = (positions < len(lookup_codes)) & (lookup_codes[safe_positions] == flat_codes)
                overlay_array.reshape(-1, 4)[matched] = lookup_colors[safe_positions[matched]]
        else:
            grayscale = np.asarray(mask.convert("L"))
            selected = np.isin(grayscale, list(foreground_values)) if foreground_values else grayscale > 0
            overlay_array[selected] = (0, 215, 255, 255)
    return image, Image.fromarray(overlay_array)


def compose_preview(
    image_path: str,
    mask_path: str,
    show_mask: bool,
    opacity: int,
    foreground_values: set[int] | None = None,
    mask_schema: dict[str, Any] | None = None,
) -> Image.Image:
    image, overlay = compose_preview_layers(
        image_path, mask_path, foreground_values, mask_schema
    )
    if not show_mask:
        return image
    alpha_scale = max(0.0, min(1.0, opacity / 100.0))
    overlay.putalpha(overlay.getchannel("A").point(lambda value: round(value * alpha_scale)))
    return Image.alpha_composite(image.convert("RGBA"), overlay)


class PreviewRenderSignals(QObject):
    finished = Signal(int, str, str, object, object, str)


class PreviewRenderTask(QRunnable):
    def __init__(
        self,
        request_id: int,
        image_path: str,
        mask_path: str,
        foreground_values: set[int] | None,
        mask_schema: dict[str, Any] | None,
        visible_label_keys: set[str] | None,
    ) -> None:
        super().__init__()
        self.setAutoDelete(False)
        self.request_id = request_id
        self.image_path = image_path
        self.mask_path = mask_path
        self.foreground_values = foreground_values
        self.mask_schema = mask_schema
        self.visible_label_keys = visible_label_keys
        self.signals = PreviewRenderSignals()

    @Slot()
    def run(self) -> None:
        try:
            image, overlay = compose_preview_layers(
                self.image_path,
                self.mask_path,
                self.foreground_values,
                self.mask_schema,
                self.visible_label_keys,
            )
            base_qimage = pil_to_qimage(image)
            overlay_qimage = pil_to_qimage(overlay)
            error = ""
        except Exception as exc:
            base_qimage = QImage()
            overlay_qimage = QImage()
            error = str(exc)
        self.signals.finished.emit(
            self.request_id,
            self.image_path,
            self.mask_path,
            base_qimage,
            overlay_qimage,
            error,
        )


class ImageCanvas(QGraphicsView):
    def __init__(self) -> None:
        super().__init__()
        self.canvas_scene = QGraphicsScene(self)
        self.setScene(self.canvas_scene)
        self.base_item = None
        self.overlay_item = None
        self._fit_mode = True
        self.setMinimumSize(400, 400)
        self.setDragMode(QGraphicsView.DragMode.ScrollHandDrag)
        self.setTransformationAnchor(QGraphicsView.ViewportAnchor.AnchorUnderMouse)
        self.setResizeAnchor(QGraphicsView.ViewportAnchor.AnchorViewCenter)
        self.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform, True)
        self.set_error("No image")

    def set_pixmap(self, pixmap: QPixmap) -> None:
        self.set_layers(pixmap, QPixmap())

    def set_layers(self, base: QPixmap, overlay: QPixmap) -> None:
        self.canvas_scene.clear()
        self.base_item = self.canvas_scene.addPixmap(base)
        self.overlay_item = self.canvas_scene.addPixmap(overlay)
        self.overlay_item.setZValue(1)
        self.canvas_scene.setSceneRect(self.base_item.boundingRect())
        self._fit_mode = True
        self.fit()

    def set_loading(self) -> None:
        self._show_message("正在加载图像...")

    def set_error(self, text: str) -> None:
        self._show_message(text)

    def _show_message(self, text: str) -> None:
        self.canvas_scene.clear()
        self.base_item = None
        self.overlay_item = None
        message = self.canvas_scene.addText(text)
        message.setDefaultTextColor(QColor("#d7dde3"))
        self.canvas_scene.setSceneRect(message.boundingRect())
        self.resetTransform()

    def set_overlay_visible(self, visible: bool) -> None:
        if self.overlay_item is not None:
            self.overlay_item.setVisible(visible)

    def set_overlay_opacity(self, opacity: int) -> None:
        if self.overlay_item is not None:
            self.overlay_item.setOpacity(max(0.0, min(1.0, opacity / 100.0)))

    def wheelEvent(self, event) -> None:  # type: ignore[no-untyped-def]
        if self.base_item is None:
            return super().wheelEvent(event)
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            factor = 1.15 if event.angleDelta().y() > 0 else 1 / 1.15
            next_scale = self.transform().m11() * factor
            if 0.03 <= next_scale <= 20.0:
                self.scale(factor, factor)
                self._fit_mode = False
            event.accept()
            return
        super().wheelEvent(event)

    def fit(self) -> None:
        if self.base_item is None:
            return
        self.resetTransform()
        self.fitInView(self.canvas_scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)
        self._fit_mode = True

    def actual_size(self) -> None:
        if self.base_item is None:
            return
        self.resetTransform()
        self._fit_mode = False

    def resizeEvent(self, event) -> None:  # type: ignore[no-untyped-def]
        super().resizeEvent(event)
        if self._fit_mode:
            self.fit()


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Scene Review Tool")
        self.resize(1500, 900)
        self.taxonomy_path: Path | None = None
        self.taxonomy = Taxonomy()
        self.workspace_dir: Path | None = None
        self.last_export_directory: Path | None = None
        self.db: ReviewDatabase | None = None
        self.dataset_id: int | None = None
        self.mask_foreground_values: set[int] | None = {1}
        self.mask_schema: dict[str, Any] | None = None
        self.visible_mask_label_keys: set[str] | None = None
        self.render_request_id = 0
        self.render_tasks: dict[int, PreviewRenderTask] = {}
        self.render_pool = QThreadPool(self)
        self.render_pool.setMaxThreadCount(1)
        self.items: list[tuple[Sample, Review]] = []
        self.current_index = -1
        self.loading = False

        self.note_timer = QTimer(self)
        self.note_timer.setSingleShot(True)
        self.note_timer.setInterval(500)
        self.note_timer.timeout.connect(self.save_quality_note)

        self.stack = QStackedWidget()
        self.project_page = self._build_project_page()
        self.review_page = self._build_review_page()
        self.stack.addWidget(self.project_page)
        self.stack.addWidget(self.review_page)
        self.setCentralWidget(self.stack)
        self._apply_style()

    def _build_project_page(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("遥感图像-mask 场景审查工具")
        title.setObjectName("title")
        subtitle = QLabel("导入数据集后逐图查看原图与 mask 叠加，人工确认质量和场景，所有结果自动保存。")
        subtitle.setObjectName("subtitle")
        layout.addWidget(title)
        layout.addWidget(subtitle)

        form_box = QGroupBox("项目导入")
        form = QFormLayout(form_box)
        self.dataset_name_edit = QLineEdit()
        self.dataset_name_edit.setPlaceholderText("例如 Greenland-512")
        self.workspace_edit = QLineEdit()
        self.workspace_edit.setPlaceholderText("请选择或输入用于保存审查进度的文件夹")
        self.taxonomy_edit = QLineEdit()
        self.taxonomy_edit.setPlaceholderText("可选；留空时使用数据集自定义场景")

        self.import_type_combo = QComboBox()
        self.import_type_combo.addItem("通用图像-Mask文件夹", "generic")
        self.import_type_combo.addItem("VOC语义分割数据集", "voc")

        self.import_options_stack = QStackedWidget()
        generic_page = QWidget()
        generic_form = QFormLayout(generic_page)
        generic_form.setContentsMargins(0, 4, 0, 4)
        self.image_root_edit = QLineEdit()
        self.image_root_edit.setPlaceholderText("请选择图像所在文件夹")
        self.mask_root_edit = QLineEdit()
        self.mask_root_edit.setPlaceholderText("请选择 Mask 所在文件夹")
        self.generic_pairing_combo = QComboBox()
        self.generic_pairing_combo.addItem("同名文件（忽略扩展名）", "same_stem")
        self.generic_pairing_combo.addItem("相对路径与文件名均相同", "relative_path_stem")
        self.generic_pairing_combo.addItem(
            "DeepGlobe（_sat 原图 / _mask 标签）", "deepglobe_suffix"
        )
        self.mask_name_prefix_edit = QLineEdit()
        self.mask_name_prefix_edit.setPlaceholderText("可选；例如 mask_")
        self.mask_name_suffix_edit = QLineEdit()
        self.mask_name_suffix_edit.setPlaceholderText("可选；例如 _instance_color_RGB")
        self.recursive_checkbox = QCheckBox("扫描子文件夹")
        self.recursive_checkbox.setChecked(True)
        generic_form.addRow("图像文件夹", self._path_row(self.image_root_edit, True))
        generic_form.addRow("Mask文件夹", self._path_row(self.mask_root_edit, True))
        generic_form.addRow("配对规则", self.generic_pairing_combo)
        generic_form.addRow("Mask文件名前缀", self.mask_name_prefix_edit)
        generic_form.addRow("Mask文件名后缀", self.mask_name_suffix_edit)
        generic_form.addRow("", self.recursive_checkbox)
        self.import_options_stack.addWidget(generic_page)

        voc_page = QWidget()
        voc_form = QFormLayout(voc_page)
        voc_form.setContentsMargins(0, 4, 0, 4)
        self.voc_root_edit = QLineEdit()
        self.voc_root_edit.setPlaceholderText("请选择 VOC 数据集根目录")
        self.image_dir_name_edit = QLineEdit()
        self.image_dir_name_edit.setPlaceholderText("默认 JPEGImages")
        self.mask_dir_name_edit = QLineEdit()
        self.mask_dir_name_edit.setPlaceholderText("默认 SegmentationClass")
        self.split_combo = QComboBox()
        self.split_combo.setEditable(True)
        self.split_combo.addItems(["all", "train", "val", "test"])
        voc_form.addRow("VOC根目录", self._path_row(self.voc_root_edit, True))
        voc_form.addRow("数据划分", self.split_combo)
        voc_form.addRow("图像目录名", self.image_dir_name_edit)
        voc_form.addRow("Mask目录名", self.mask_dir_name_edit)
        self.import_options_stack.addWidget(voc_page)
        self.import_type_combo.currentIndexChanged.connect(self.import_options_stack.setCurrentIndex)

        self.mask_values_edit = QLineEdit()
        self.mask_values_edit.setPlaceholderText("可选；兼容旧式灰度前景值，例如 1,2")
        self.mask_name_edit = QLineEdit()
        self.mask_name_edit.setPlaceholderText("可选，例如 green space；默认 foreground")
        form.addRow("数据集名称", self.dataset_name_edit)
        form.addRow("导入方式", self.import_type_combo)
        form.addRow(self.import_options_stack)
        form.addRow("工作区", self._path_row(self.workspace_edit, True))
        form.addRow("场景体系 JSON（可选）", self._path_row(self.taxonomy_edit, False))
        form.addRow("手动前景值（可选）", self.mask_values_edit)
        form.addRow("默认前景名称", self.mask_name_edit)
        layout.addWidget(form_box)

        class_box = QGroupBox("Mask 类别映射")
        class_layout = QVBoxLayout(class_box)
        class_hint = QLabel(
            "扫描后自动识别标签。背景只是候选项；全分类 Mask 可将所有有效标签设为 class。"
        )
        class_hint.setObjectName("hint")
        class_hint.setWordWrap(True)
        self.mask_class_table = QTableWidget(0, 4)
        self.mask_class_table.setHorizontalHeaderLabels(["颜色", "原始标签", "角色", "类别名称"])
        header = self.mask_class_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.mask_class_table.setMinimumHeight(170)
        class_layout.addWidget(class_hint)
        class_layout.addWidget(self.mask_class_table)
        layout.addWidget(class_box)

        action_row = QHBoxLayout()
        scan_btn = QPushButton("扫描预览")
        scan_btn.clicked.connect(self.preview_import)
        create_btn = QPushButton("导入并开始审查")
        create_btn.clicked.connect(self.create_or_open_project)
        open_btn = QPushButton("打开已有工作区")
        open_btn.clicked.connect(self.open_existing_workspace)
        action_row.addWidget(scan_btn)
        action_row.addWidget(create_btn)
        action_row.addWidget(open_btn)
        action_row.addStretch()
        layout.addLayout(action_row)

        preview_row = QHBoxLayout()
        self.project_message = QLabel("请选择导入方式和数据路径，然后先执行扫描预览。")
        self.project_message.setWordWrap(True)
        self.project_message.setAlignment(Qt.AlignmentFlag.AlignTop)
        self.preview_image_label = QLabel("尚未生成配对预览")
        self.preview_image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_image_label.setMinimumSize(360, 220)
        self.preview_image_label.setMaximumSize(520, 300)
        self.preview_image_label.setFrameShape(QFrame.Shape.StyledPanel)
        preview_row.addWidget(self.project_message, 1)
        preview_row.addWidget(self.preview_image_label, 1)
        layout.addLayout(preview_row)
        layout.addStretch()
        return page

    def populate_mask_class_table(self, schema: dict[str, Any]) -> None:
        labels = schema.get("labels", [])
        self.mask_class_table.setRowCount(len(labels))
        for row, label in enumerate(labels):
            color = str(label.get("color") or "#00d7ff")
            color_item = QTableWidgetItem(color)
            color_item.setBackground(QColor(color))
            color_item.setFlags(color_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            raw_item = QTableWidgetItem(mask_label_text(label.get("value", "")))
            raw_item.setData(Qt.ItemDataRole.UserRole, dict(label))
            raw_item.setFlags(raw_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            role_combo = QComboBox()
            role_combo.addItems(["background", "class", "ignore"])
            role_combo.setCurrentText(str(label.get("role") or "class"))
            name_item = QTableWidgetItem(str(label.get("name") or ""))
            self.mask_class_table.setItem(row, 0, color_item)
            self.mask_class_table.setItem(row, 1, raw_item)
            self.mask_class_table.setCellWidget(row, 2, role_combo)
            self.mask_class_table.setItem(row, 3, name_item)

    def mask_schema_from_table(self, encoding: str) -> dict[str, Any] | None:
        if self.mask_class_table.rowCount() == 0:
            return None
        labels: list[dict[str, Any]] = []
        for row in range(self.mask_class_table.rowCount()):
            raw_item = self.mask_class_table.item(row, 1)
            role_combo = self.mask_class_table.cellWidget(row, 2)
            name_item = self.mask_class_table.item(row, 3)
            if raw_item is None or not isinstance(role_combo, QComboBox):
                continue
            label = raw_item.data(Qt.ItemDataRole.UserRole)
            if not isinstance(label, dict):
                continue
            labels.append(
                {
                    "key": label["key"],
                    "value": label["value"],
                    "role": role_combo.currentText(),
                    "name": name_item.text().strip() if name_item else "",
                    "color": label["color"],
                }
            )
        if not labels:
            return None
        has_background = any(label["role"] == "background" for label in labels)
        return {
            "schema_version": 3,
            "encoding": encoding,
            "background_mode": "explicit" if has_background else "none",
            "labels": labels,
        }

    def _path_row(self, edit: QLineEdit, directory: bool) -> QWidget:
        row = QWidget()
        layout = QHBoxLayout(row)
        layout.setContentsMargins(0, 0, 0, 0)
        btn = QPushButton("选择")
        btn.clicked.connect(lambda: self._browse_path(edit, directory))
        layout.addWidget(edit)
        layout.addWidget(btn)
        return row

    def _browse_path(self, edit: QLineEdit, directory: bool) -> None:
        if directory:
            value = QFileDialog.getExistingDirectory(self, "选择目录", edit.text())
        else:
            value, _ = QFileDialog.getOpenFileName(self, "选择文件", edit.text(), "JSON (*.json);;All files (*.*)")
        if value:
            edit.setText(value)

    def _build_review_page(self) -> QWidget:
        page = QWidget()
        root = QVBoxLayout(page)

        toolbar = QHBoxLayout()
        self.dataset_label = QLabel("")
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("搜索文件名 / 样本 ID / 备注")
        self.search_edit.returnPressed.connect(self.reload_samples)
        self.filter_combo = QComboBox()
        self.filter_combo.addItems(["all", "unreviewed", "accepted", "needs_correction", "rejected", "unassigned"])
        self.filter_combo.currentTextChanged.connect(self.reload_samples)
        stats_btn = QPushButton("导出统计")
        stats_btn.clicked.connect(self.export_stats)
        quality_import_btn = QPushButton("导入质量进度")
        quality_import_btn.clicked.connect(self.import_quality_progress)
        quality_export_btn = QPushButton("按质量导出")
        quality_export_btn.clicked.connect(self.export_quality)
        self.selected_export_btn = QPushButton("导出选中（0）")
        self.selected_export_btn.setEnabled(False)
        self.selected_export_btn.clicked.connect(self.export_selected)
        accepted_export_btn = QPushButton("导出合格数据")
        accepted_export_btn.clicked.connect(self.export_accepted)
        export_btn = QPushButton("按场景导出")
        export_btn.clicked.connect(self.export_by_scene)
        back_btn = QPushButton("项目页")
        back_btn.clicked.connect(lambda: self.stack.setCurrentWidget(self.project_page))
        toolbar.addWidget(self.dataset_label)
        toolbar.addWidget(self.search_edit, 2)
        toolbar.addWidget(self.filter_combo)
        toolbar.addWidget(stats_btn)
        toolbar.addWidget(quality_import_btn)
        toolbar.addWidget(quality_export_btn)
        toolbar.addWidget(self.selected_export_btn)
        toolbar.addWidget(accepted_export_btn)
        toolbar.addWidget(export_btn)
        toolbar.addWidget(back_btn)
        root.addLayout(toolbar)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        self.sample_list = QListWidget()
        self.sample_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.sample_list.setToolTip("Ctrl+单击选择多个样本；Shift+单击选择连续区间")
        self.sample_list.currentRowChanged.connect(self.on_sample_selected)
        self.sample_list.itemSelectionChanged.connect(self.update_selected_export_button)
        splitter.addWidget(self.sample_list)

        center = QWidget()
        center_layout = QVBoxLayout(center)
        view_bar = QHBoxLayout()
        self.mask_checkbox = QCheckBox("显示 mask")
        self.mask_checkbox.setChecked(True)
        self.mask_checkbox.stateChanged.connect(self.update_mask_visibility)
        self.opacity_slider = QSlider(Qt.Orientation.Horizontal)
        self.opacity_slider.setRange(5, 95)
        self.opacity_slider.setValue(45)
        self.opacity_slider.valueChanged.connect(self.update_mask_opacity)
        fit_btn = QPushButton("适合窗口")
        fit_btn.clicked.connect(lambda: self.image_canvas.fit())
        actual_btn = QPushButton("100%")
        actual_btn.clicked.connect(lambda: self.image_canvas.actual_size())
        view_bar.addWidget(self.mask_checkbox)
        view_bar.addWidget(QLabel("透明度"))
        view_bar.addWidget(self.opacity_slider)
        view_bar.addWidget(fit_btn)
        view_bar.addWidget(actual_btn)
        center_layout.addLayout(view_bar)

        image_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.image_canvas = ImageCanvas()
        image_splitter.addWidget(self.image_canvas)
        legend_box = QGroupBox("Mask 图例")
        legend_layout = QVBoxLayout(legend_box)
        self.mask_legend = QTreeWidget()
        self.mask_legend.setHeaderLabels(["颜色", "类别名称", "原始标签", "角色"])
        self.mask_legend.setRootIsDecorated(False)
        self.mask_legend.setAlternatingRowColors(True)
        self.mask_legend.itemChanged.connect(self.on_mask_legend_changed)
        legend_layout.addWidget(self.mask_legend)
        legend_box.setMinimumWidth(230)
        image_splitter.addWidget(legend_box)
        image_splitter.setSizes([760, 240])
        center_layout.addWidget(image_splitter, 1)
        self.path_label = QLabel("")
        self.path_label.setWordWrap(True)
        center_layout.addWidget(self.path_label)
        splitter.addWidget(center)

        panel = self._build_review_panel()
        splitter.addWidget(panel)
        splitter.setSizes([300, 820, 380])
        root.addWidget(splitter, 1)

        bottom = QHBoxLayout()
        prev_btn = QPushButton("上一张")
        next_btn = QPushButton("下一张")
        save_btn = QPushButton("保存")
        next_unreviewed_btn = QPushButton("下一张未审")
        prev_btn.clicked.connect(self.prev_sample)
        next_btn.clicked.connect(self.next_sample)
        save_btn.clicked.connect(self.save_current_review)
        next_unreviewed_btn.clicked.connect(self.next_unreviewed)
        self.status_label = QLabel("")
        bottom.addWidget(prev_btn)
        bottom.addWidget(next_btn)
        bottom.addWidget(next_unreviewed_btn)
        bottom.addWidget(save_btn)
        bottom.addWidget(self.status_label, 1)
        root.addLayout(bottom)
        return page

    def _build_review_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)

        quality_box = QGroupBox("质量状态")
        q_layout = QVBoxLayout(quality_box)
        self.quality_group = QButtonGroup(self)
        for status in QUALITY_STATUSES:
            radio = QRadioButton(status)
            self.quality_group.addButton(radio)
            q_layout.addWidget(radio)
            if status == "unreviewed":
                radio.setChecked(True)
        self.quality_group.buttonClicked.connect(self.save_quality_status_only)
        self.note_edit = QTextEdit()
        self.note_edit.setAcceptRichText(False)
        self.note_edit.setMaximumHeight(120)
        self.note_edit.textChanged.connect(self.on_quality_note_changed)
        q_layout.addWidget(QLabel("质量备注"))
        q_layout.addWidget(self.note_edit)
        self.note_save_label = QLabel("")
        q_layout.addWidget(self.note_save_label)
        layout.addWidget(quality_box)

        common_box = QGroupBox("当前数据集常用场景")
        common_layout = QVBoxLayout(common_box)
        self.common_scene_list = QListWidget()
        self.common_scene_list.itemClicked.connect(self.apply_common_scene)
        common_layout.addWidget(self.common_scene_list)
        layout.addWidget(common_box, 2)

        scene_box = QGroupBox("场景选择")
        scene_layout = QFormLayout(scene_box)
        self.scene_source_combo = QComboBox()
        self.scene_source_combo.addItems(SCENE_SOURCES)
        self.scene_source_combo.currentTextChanged.connect(self.on_scene_source_changed)
        self.domain_combo = QComboBox()
        self.domain_combo.currentIndexChanged.connect(self.reload_scene_combo)
        self.scene_combo = QComboBox()
        self.custom_en_edit = QLineEdit()
        self.custom_en_edit.setPlaceholderText("例如 urban waterfront")
        self.custom_zh_edit = QLineEdit()
        self.custom_zh_edit.setPlaceholderText("中文说明")
        self.mapping_edit = QLineEdit()
        self.mapping_edit.setPlaceholderText("可选：映射到正式场景，如 river")
        self.scene_status_combo = QComboBox()
        self.scene_status_combo.addItems(["unassigned", "assigned", "uncertain"])
        scene_layout.addRow("来源", self.scene_source_combo)
        scene_layout.addRow("一级", self.domain_combo)
        scene_layout.addRow("二级", self.scene_combo)
        scene_layout.addRow("自定义英文", self.custom_en_edit)
        scene_layout.addRow("自定义中文", self.custom_zh_edit)
        scene_layout.addRow("映射建议", self.mapping_edit)
        scene_layout.addRow("场景状态", self.scene_status_combo)
        scene_actions = QWidget()
        scene_actions_layout = QHBoxLayout(scene_actions)
        scene_actions_layout.setContentsMargins(0, 0, 0, 0)
        confirm_scene_btn = QPushButton("确认场景")
        confirm_scene_btn.clicked.connect(self.confirm_scene_selection)
        self.reset_taxonomy_btn = QPushButton("改用正式体系")
        self.reset_taxonomy_btn.clicked.connect(self.reset_to_taxonomy_scene)
        clear_scene_btn = QPushButton("清空")
        clear_scene_btn.clicked.connect(self.clear_scene_selection)
        scene_actions_layout.addWidget(confirm_scene_btn)
        scene_actions_layout.addWidget(self.reset_taxonomy_btn)
        scene_actions_layout.addWidget(clear_scene_btn)
        scene_layout.addRow("", scene_actions)
        self.scene_hint_label = QLabel("点击常用场景只填入候选，确认后保存。")
        self.scene_hint_label.setObjectName("hint")
        self.scene_hint_label.setWordWrap(True)
        scene_layout.addRow("", self.scene_hint_label)
        layout.addWidget(scene_box)


        self.refresh_taxonomy_controls()
        return panel

    def _apply_style(self) -> None:
        self.setStyleSheet(
            """
            QWidget { background: #202225; color: #edf0f2; font-size: 14px; }
            QLineEdit, QTextEdit, QComboBox, QListWidget {
                background: #2c3035; border: 1px solid #474d55; border-radius: 6px;
                padding: 6px; selection-background-color: #2d7d9a;
            }
            QPushButton {
                background: #39515f; border: 1px solid #526b78; border-radius: 6px;
                padding: 7px 12px;
            }
            QPushButton:hover { background: #446575; }
            QGroupBox {
                border: 1px solid #444a52; border-radius: 8px; margin-top: 14px; padding: 10px;
                font-weight: 600;
            }
            QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 4px; }
            QLabel#title { font-size: 24px; font-weight: 700; }
            QLabel#subtitle { color: #b9c0c7; }
            QLabel#hint { color: #aeb7bf; font-size: 12px; }
            QListWidget::item { padding: 6px; }
            QListWidget::item:selected { background: #2d7d9a; }
            """
        )

    def collect_import_data(self) -> dict[str, Any]:
        dataset_name = self.dataset_name_edit.text().strip()
        workspace_text = self.workspace_edit.text().strip()
        if not dataset_name:
            raise ValueError("请填写数据集名称。")
        if not workspace_text:
            raise ValueError("请选择或输入工作区。")
        workspace = Path(workspace_text)
        taxonomy_text = self.taxonomy_edit.text().strip()
        taxonomy_path: Path | None = None
        taxonomy = Taxonomy()
        if taxonomy_text:
            taxonomy_path = Path(taxonomy_text)
            if not taxonomy_path.is_file():
                raise ValueError("场景体系 JSON 不存在；不使用场景体系时请将该项留空。")
            taxonomy = Taxonomy(taxonomy_path)
            if not taxonomy.scenes:
                raise ValueError("场景体系 JSON 中没有可用的场景。")

        import_type = self.import_type_combo.currentData()
        if import_type == "voc":
            if not self.voc_root_edit.text().strip():
                raise ValueError("请选择VOC根目录。")
            dataset_root = Path(self.voc_root_edit.text().strip())
            image_root = dataset_root
            mask_root = dataset_root
            mode = "voc_segmentation"
            image_dir_name = self.image_dir_name_edit.text().strip() or "JPEGImages"
            mask_dir_name = self.mask_dir_name_edit.text().strip() or "SegmentationClass"
            split = self.split_combo.currentText().strip() or "all"
            recursive = False
            mask_name_prefix = ""
            mask_name_suffix = ""
            image_scan_root = dataset_root / image_dir_name
            mask_scan_root = dataset_root / mask_dir_name
            if not dataset_root.is_dir() or not image_scan_root.is_dir() or not mask_scan_root.is_dir():
                raise ValueError("VOC根目录中没有找到图像目录或Mask目录。")
        else:
            if not self.image_root_edit.text().strip() or not self.mask_root_edit.text().strip():
                raise ValueError("请选择图像文件夹和Mask文件夹。")
            image_root = Path(self.image_root_edit.text().strip())
            mask_root = Path(self.mask_root_edit.text().strip())
            mode = str(self.generic_pairing_combo.currentData())
            image_dir_name = ""
            mask_dir_name = ""
            split = "all"
            recursive = self.recursive_checkbox.isChecked()
            mask_name_prefix = self.mask_name_prefix_edit.text().strip()
            mask_name_suffix = self.mask_name_suffix_edit.text().strip()
            image_scan_root = image_root
            mask_scan_root = mask_root
            if not image_root.is_dir() or not mask_root.is_dir():
                raise ValueError("请检查图像文件夹和Mask文件夹。")

        try:
            manual_foreground_values = {
                int(value.strip())
                for value in self.mask_values_edit.text().split(",")
                if value.strip()
            }
        except ValueError as exc:
            raise ValueError("Mask前景标签值必须是整数，多个值使用英文逗号分隔。") from exc

        samples, warnings = scan_dataset(
            dataset_name,
            image_root,
            mask_root,
            mode,
            image_dir_name,
            mask_dir_name,
            split,
            recursive,
            mask_name_prefix,
            mask_name_suffix,
        )
        if not samples:
            raise ValueError("没有找到成功配对的image-mask样本。请检查路径和配对规则。")
        inspection = inspect_sample_pairs(samples)
        if not inspection["mask_labels"]:
            raise ValueError("抽样 Mask 中没有检测到离散标签；请确认标签图不是连续色彩图或压缩图。")
        if inspection["label_overflow"]:
            warnings.append(
                f"{inspection['label_overflow']} 张抽样 Mask 的颜色超过 4096 种，可能不是离散标签图。"
            )

        schema = infer_mask_schema(inspection, self.mask_name_edit.text())
        table_schema = self.mask_schema_from_table(inspection["mask_encoding"])
        detected_keys = {label["key"] for label in inspection["mask_labels"]}
        if table_schema and {label["key"] for label in table_schema["labels"]} == detected_keys:
            schema = table_schema
        if manual_foreground_values and inspection["mask_encoding"] in {"indexed", "mixed"}:
            for label in schema["labels"]:
                value = label["value"]
                if isinstance(value, int):
                    if value in manual_foreground_values:
                        label["role"] = "class"
                        if len(manual_foreground_values) == 1:
                            label["name"] = self.mask_name_edit.text().strip() or "foreground"
                    elif value == 0:
                        label["role"] = "background"
                        label["name"] = "background"
                    else:
                        label["role"] = "ignore"
        has_background = any(label["role"] == "background" for label in schema["labels"])
        schema["background_mode"] = "explicit" if has_background else "none"
        validate_mask_schema(schema)

        foreground_values = {
            int(label["value"])
            for label in schema["labels"]
            if label["role"] == "class" and isinstance(label["value"], int)
        }
        background_values = [
            label["value"] for label in schema["labels"] if label["role"] == "background"
        ]
        class_names = [label["name"] for label in schema["labels"] if label["role"] == "class"]

        image_count = len(supported_files(image_scan_root, IMAGE_EXTS, recursive))
        mask_count = len(supported_files(mask_scan_root, MASK_EXTS, recursive))
        deepglobe_check: dict[str, int] | None = None
        if mode == "deepglobe_suffix":
            image_count = len(
                deepglobe_role_files(
                    image_scan_root, IMAGE_EXTS, DEEPGLOBE_IMAGE_SUFFIX, recursive
                )
            )
            mask_count = len(
                deepglobe_role_files(
                    mask_scan_root, MASK_EXTS, DEEPGLOBE_MASK_SUFFIX, recursive
                )
            )
            deepglobe_check = {
                "sat_images": image_count,
                "mask_labels": mask_count,
                "paired": len(samples),
            }

        config = {
            "dataset_name": dataset_name,
            "import_type": import_type,
            "image_root": normalized_path(image_root),
            "mask_root": normalized_path(mask_root),
            "taxonomy_path": normalized_path(taxonomy_path) if taxonomy_path else "",
            "pairing": {
                "mode": mode,
                "image_dir_name": image_dir_name,
                "mask_dir_name": mask_dir_name,
                "split": split,
                "recursive": recursive,
                "mask_name_prefix": mask_name_prefix,
                "mask_name_suffix": mask_name_suffix,
            },
            "mask": {
                "type": "binary" if len(class_names) <= 1 else "multiclass",
                "schema_version": schema["schema_version"],
                "encoding": schema["encoding"],
                "background_mode": schema["background_mode"],
                "labels": schema["labels"],
                "background_value": background_values[0] if background_values else None,
                "background_values": background_values,
                "foreground_values": sorted(foreground_values),
                "foreground_name": self.mask_name_edit.text().strip() or "foreground",
            },
        }
        return {
            "config": config,
            "samples": samples,
            "warnings": warnings,
            "inspection": inspection,
            "mask_schema": schema,
            "image_count": image_count,
            "mask_count": mask_count,
            "deepglobe_check": deepglobe_check,
            "image_root": image_root,
            "mask_root": mask_root,
            "workspace": workspace,
            "taxonomy_path": taxonomy_path,
            "taxonomy": taxonomy,
            "foreground_values": foreground_values,
        }

    def format_import_summary(self, data: dict[str, Any]) -> str:
        inspection = data["inspection"]
        labels = inspection["mask_labels"]
        value_text = ", ".join(mask_label_text(label["value"]) for label in labels) if labels else "未能枚举"
        lines = [
            f"目录图像：{data['image_count']}    目录Mask：{data['mask_count']}",
            f"成功配对：{len(data['samples'])}    警告：{len(data['warnings'])}",
            (
                f"抽样检查：{inspection['checked']} 对    "
                f"Mask类型：{inspection['mask_encoding']} {inspection['mask_modes']}    "
                f"标签：{value_text}"
            ),
            f"抽样尺寸不一致：{inspection['size_mismatches']}    读取错误：{inspection['read_errors']}",
            (
                f"场景模式：已加载体系（{len(data['taxonomy'].scenes)} 个场景）"
                if data["taxonomy"].scenes
                else "场景模式：数据集自定义场景（未使用场景体系 JSON）"
            ),
        ]
        deepglobe_check = data.get("deepglobe_check")
        if deepglobe_check:
            lines.insert(
                1,
                (
                    "DeepGlobe 自动检查："
                    f"_sat 原图 {deepglobe_check['sat_images']}    "
                    f"_mask 标签 {deepglobe_check['mask_labels']}    "
                    f"成功一对一配对 {deepglobe_check['paired']}"
                ),
            )
        class_count = sum(label["role"] == "class" for label in data["mask_schema"]["labels"])
        ignored_count = sum(label["role"] == "ignore" for label in data["mask_schema"]["labels"])
        background_count = sum(
            label["role"] == "background" for label in data["mask_schema"]["labels"]
        )
        background_text = f"{background_count} 个背景标签" if background_count else "无背景（全分类）"
        lines.append(f"类别映射：{class_count} 个有效类别，{background_text}，{ignored_count} 个忽略标签")
        if data["warnings"]:
            lines.append("部分警告：")
            lines.extend(f"- {warning}" for warning in data["warnings"][:4])
            if len(data["warnings"]) > 4:
                lines.append(f"- 其余 {len(data['warnings']) - 4} 条未显示")
        return "\n".join(lines)

    def show_import_preview_image(self, data: dict[str, Any]) -> None:
        sample = data["samples"][0]
        image = compose_preview(
            sample.image_path,
            sample.mask_path,
            True,
            45,
            data["foreground_values"],
            data["mask_schema"],
        )
        pixmap = pil_to_qpixmap(image).scaled(
            self.preview_image_label.size(),
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        self.preview_image_label.setPixmap(pixmap)
        self.preview_image_label.setToolTip(
            f"image: {sample.image_path}\nmask: {sample.mask_path}"
        )

    def preview_import(self) -> None:
        try:
            data = self.collect_import_data()
            self.project_message.setText(self.format_import_summary(data))
            self.populate_mask_class_table(data["mask_schema"])
            self.show_import_preview_image(data)
        except Exception as exc:
            self.project_message.setText(f"扫描失败：{exc}")
            self.preview_image_label.setPixmap(QPixmap())
            self.preview_image_label.setText("无法生成预览")
            QMessageBox.warning(self, "扫描失败", str(exc))

    def create_or_open_project(self) -> None:
        if not self.save_quality_note():
            return
        try:
            data = self.collect_import_data()
            config = data["config"]
            samples = data["samples"]
            warnings = data["warnings"]
            image_root = data["image_root"]
            mask_root = data["mask_root"]
            workspace = data["workspace"]
            taxonomy_path = data["taxonomy_path"]
            taxonomy = data["taxonomy"]
            foreground_values = data["foreground_values"]
            mask_schema = data["mask_schema"]
            dataset_name = config["dataset_name"]

            if (workspace / "review.sqlite3").exists():
                raise FileExistsError(
                    "该工作区已经存在审查数据库。继续以前的工作请使用“打开已有工作区”；"
                    "导入新数据集请更换工作区。"
                )
            summary = self.format_import_summary(data)
            answer = QMessageBox.question(
                self,
                "确认导入",
                f"{summary}\n\n确认建立审查项目吗？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return

            self.taxonomy_path = taxonomy_path
            self.taxonomy = taxonomy
            self.refresh_taxonomy_controls()
            self.mask_foreground_values = foreground_values
            self.mask_schema = mask_schema
            self.refresh_mask_legend()
            workspace.mkdir(parents=True, exist_ok=True)
            self.workspace_dir = workspace
            self.workspace_edit.setText(str(workspace))
            (workspace / "exports").mkdir(exist_ok=True)
            (workspace / "cache").mkdir(exist_ok=True)
            (workspace / "project.json").write_text(json.dumps(config, indent=2, ensure_ascii=False), encoding="utf-8")
            self.current_index = -1
            self.items = []
            if self.db:
                self.db.close()
            self.db = ReviewDatabase(workspace / "review.sqlite3")
            self.dataset_id = self.db.create_dataset(dataset_name, image_root, mask_root, config, self.taxonomy)
            self.db.add_samples(samples, self.dataset_id)
            class_names = [
                label["name"] for label in mask_schema["labels"] if label["role"] == "class"
            ]
            msg = (
                f"导入完成：成功配对 {len(samples)} 个样本。"
                f"Mask 类型：{mask_schema['encoding']}；有效类别：{', '.join(class_names) or '无'}。"
            )
            if warnings:
                msg += f" 未配对或警告 {len(warnings)} 条，已跳过。"
            self.project_message.setText(msg)
            self.load_review_page()
        except Exception as exc:
            QMessageBox.critical(self, "导入失败", str(exc))

    def open_existing_workspace(self) -> None:
        path = QFileDialog.getExistingDirectory(self, "选择已有工作区", self.workspace_edit.text())
        if not path:
            return
        workspace = Path(path)
        db_path = workspace / "review.sqlite3"
        if not db_path.exists():
            QMessageBox.warning(self, "无法打开", "该工作区没有 review.sqlite3。")
            return
        if not self.save_quality_note():
            return
        self.current_index = -1
        self.items = []
        if self.db:
            self.db.close()
        self.db = ReviewDatabase(db_path)
        self.workspace_dir = workspace
        self.workspace_edit.setText(str(workspace))
        self.dataset_id = self.db.latest_dataset_id()
        if self.dataset_id is None:
            QMessageBox.warning(self, "无法打开", "数据库中没有数据集。")
            return
        row = self.db.latest_dataset()
        self.taxonomy_path = None
        self.taxonomy = Taxonomy()
        self.mask_schema = None
        if row and row["config_json"]:
            config = json.loads(row["config_json"])
            configured_taxonomy = str(config.get("taxonomy_path") or "").strip()
            selected_taxonomy = self.taxonomy_edit.text().strip()
            for candidate_text in (configured_taxonomy, selected_taxonomy):
                if not candidate_text:
                    continue
                candidate = Path(candidate_text)
                if not candidate.is_file():
                    continue
                try:
                    loaded_taxonomy = Taxonomy(candidate)
                except (OSError, UnicodeError, json.JSONDecodeError):
                    continue
                if loaded_taxonomy.scenes:
                    self.taxonomy_path = candidate
                    self.taxonomy = loaded_taxonomy
                    break
            mask_config = config.get("mask", {})
            labels = mask_config.get("labels")
            if isinstance(labels, list) and labels:
                self.mask_schema = {
                    "schema_version": int(mask_config.get("schema_version", 2)),
                    "encoding": str(mask_config.get("encoding", "unknown")),
                    "background_mode": str(
                        mask_config.get("background_mode")
                        or ("explicit" if any(label.get("role") == "background" for label in labels) else "none")
                    ),
                    "labels": labels,
                }
            values = mask_config.get("foreground_values")
            self.mask_foreground_values = {int(value) for value in values} if values else None
        self.refresh_mask_legend()
        self.refresh_taxonomy_controls()
        self.load_review_page()

    def load_review_page(self) -> None:
        if not self.db or self.dataset_id is None:
            return
        row = self.db.latest_dataset()
        self.dataset_label.setText(f"数据集：{row['name'] if row else self.dataset_id}")
        self.reload_samples()
        self.stack.setCurrentWidget(self.review_page)

    def reload_samples(self) -> None:
        if not self.db or self.dataset_id is None:
            return
        if not self.save_quality_note():
            return
        self.current_index = -1
        self.items = self.db.samples(self.dataset_id, self.filter_combo.currentText(), self.search_edit.text())
        self.sample_list.clear()
        for index, (sample, review) in enumerate(self.items):
            item = QListWidgetItem(self.sample_list_item_text(index, sample, review))
            item.setData(Qt.ItemDataRole.UserRole, sample.id)
            self.sample_list.addItem(item)
        self.update_common_scenes()
        self.status_label.setText(self.format_stats())
        if self.items:
            self.sample_list.setCurrentRow(0)
        else:
            self.image_canvas.set_error("当前过滤条件下没有样本")

    def on_sample_selected(self, row: int) -> None:
        if row < 0 or row >= len(self.items):
            return
        previous = self.current_index
        if not self.save_quality_note():
            self.sample_list.blockSignals(True)
            self.sample_list.setCurrentRow(previous)
            self.sample_list.blockSignals(False)
            return
        self.current_index = row
        sample, review = self.items[row]
        self.loading = True
        for button in self.quality_group.buttons():
            button.setChecked(button.text() == review.quality_status)
        has_saved_scene = bool(review.primary_level2_scene or review.custom_scene_name_en)
        source = review.scene_source if review.scene_source in SCENE_SOURCES else self.default_scene_source()
        if not has_saved_scene and review.scene_status == "unassigned":
            source = self.default_scene_source()
        self.scene_source_combo.setCurrentText(source)
        if review.level1_id:
            idx = self.domain_combo.findData(review.level1_id)
            if idx < 0 and source == "taxonomy":
                self.domain_combo.addItem(
                    f"{review.level1_id} {review.level1_name}".strip(),
                    review.level1_id,
                )
                idx = self.domain_combo.count() - 1
            if idx >= 0:
                self.domain_combo.setCurrentIndex(idx)
        self.reload_scene_combo()
        if review.primary_level2_scene:
            idx = self.scene_combo.findText(review.primary_level2_scene)
            if idx < 0 and source == "taxonomy":
                self.scene_combo.addItem(review.primary_level2_scene)
                idx = self.scene_combo.count() - 1
            if idx >= 0:
                self.scene_combo.setCurrentIndex(idx)
        self.custom_en_edit.setText(review.custom_scene_name_en)
        self.custom_zh_edit.setText(review.custom_scene_name_zh)
        self.mapping_edit.setText(review.taxonomy_mapping_suggestion)
        self.scene_status_combo.setCurrentText(review.scene_status)
        self.note_edit.setPlainText(review.reviewer_note)
        self.note_save_label.setText('已保存' if review.reviewer_note else '')
        self.loading = False
        self.refresh_image()

    def refresh_mask_legend(self) -> None:
        if not hasattr(self, "mask_legend"):
            return
        self.mask_legend.blockSignals(True)
        self.mask_legend.clear()
        visible_keys: set[str] = set()
        if self.mask_schema and self.mask_schema.get("labels"):
            for label in self.mask_schema["labels"]:
                role = str(label.get("role") or "class")
                key = str(label.get("key") or "")
                color = str(label.get("color") or "#00d7ff")
                item = QTreeWidgetItem(
                    ["■", str(label.get("name") or ""), mask_label_text(label.get("value", "")), role]
                )
                item.setData(0, Qt.ItemDataRole.UserRole, key)
                item.setForeground(0, QColor(color))
                if role == "class":
                    item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                    item.setCheckState(0, Qt.CheckState.Checked)
                    visible_keys.add(key)
                else:
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsUserCheckable)
                self.mask_legend.addTopLevelItem(item)
        elif self.mask_foreground_values:
            values = ", ".join(str(value) for value in sorted(self.mask_foreground_values))
            item = QTreeWidgetItem(["■", "foreground", values, "class"])
            item.setForeground(0, QColor("#00d7ff"))
            self.mask_legend.addTopLevelItem(item)
        self.visible_mask_label_keys = visible_keys if self.mask_schema else None
        for column in range(4):
            self.mask_legend.resizeColumnToContents(column)
        self.mask_legend.blockSignals(False)

    def on_mask_legend_changed(self, _item: QTreeWidgetItem, _column: int) -> None:
        visible_keys: set[str] = set()
        for index in range(self.mask_legend.topLevelItemCount()):
            item = self.mask_legend.topLevelItem(index)
            if item.checkState(0) == Qt.CheckState.Checked:
                key = item.data(0, Qt.ItemDataRole.UserRole)
                if key:
                    visible_keys.add(str(key))
        self.visible_mask_label_keys = visible_keys
        self.refresh_image(clear_canvas=False)

    def update_mask_visibility(self, _state: int | None = None) -> None:
        self.image_canvas.set_overlay_visible(self.mask_checkbox.isChecked())

    def update_mask_opacity(self, value: int) -> None:
        self.image_canvas.set_overlay_opacity(value)

    def refresh_image(self, clear_canvas: bool = True) -> None:
        if self.current_index < 0 or self.current_index >= len(self.items):
            return
        sample, _review = self.items[self.current_index]
        for old_request_id, old_task in list(self.render_tasks.items()):
            if self.render_pool.tryTake(old_task):
                self.render_tasks.pop(old_request_id, None)
        self.render_request_id += 1
        request_id = self.render_request_id
        task = PreviewRenderTask(
            request_id,
            sample.image_path,
            sample.mask_path,
            set(self.mask_foreground_values) if self.mask_foreground_values else None,
            self.mask_schema,
            set(self.visible_mask_label_keys) if self.visible_mask_label_keys is not None else None,
        )
        task.signals.finished.connect(self.on_preview_rendered)
        self.render_tasks[request_id] = task
        if clear_canvas:
            self.image_canvas.set_loading()
        self.path_label.setText(f"正在加载：{sample.image_path}")
        self.render_pool.start(task)

    def on_preview_rendered(
        self,
        request_id: int,
        image_path: str,
        mask_path: str,
        base_qimage: QImage,
        overlay_qimage: QImage,
        error: str,
    ) -> None:
        self.render_tasks.pop(request_id, None)
        if request_id != self.render_request_id:
            return
        if error:
            self.image_canvas.set_error(error)
            self.path_label.setText(f"读取失败：{image_path}\n{error}")
            return
        self.image_canvas.set_layers(QPixmap.fromImage(base_qimage), QPixmap.fromImage(overlay_qimage))
        self.update_mask_visibility()
        self.update_mask_opacity(self.opacity_slider.value())
        self.path_label.setText(f"image: {image_path}\nmask: {mask_path}")

    def reload_scene_combo(self) -> None:
        if not hasattr(self, "scene_combo"):
            return
        domain_id = self.domain_combo.currentData()
        self.scene_combo.clear()
        self.scene_combo.addItem("")
        for scene in self.taxonomy.by_domain.get(domain_id, []):
            self.scene_combo.addItem(scene.scene)

    def default_scene_source(self) -> str:
        return "taxonomy" if self.taxonomy.scenes else "dataset_custom"

    def refresh_taxonomy_controls(self) -> None:
        if not hasattr(self, "domain_combo"):
            return
        self.domain_combo.blockSignals(True)
        self.domain_combo.clear()
        for domain_id, name_zh, name_en in self.taxonomy.domains:
            self.domain_combo.addItem(f"{domain_id} {name_zh} / {name_en}", domain_id)
        self.domain_combo.blockSignals(False)
        taxonomy_available = bool(self.taxonomy.scenes)
        model = self.scene_source_combo.model()
        taxonomy_index = self.scene_source_combo.findText("taxonomy")
        if taxonomy_index >= 0 and hasattr(model, "item"):
            taxonomy_item = model.item(taxonomy_index)
            if taxonomy_item is not None:
                taxonomy_item.setEnabled(taxonomy_available)
        self.reset_taxonomy_btn.setEnabled(taxonomy_available)
        source = self.default_scene_source()
        self.scene_source_combo.setCurrentText(source)
        self.reload_scene_combo()
        self.on_scene_source_changed(source)
        if taxonomy_available:
            self.scene_hint_label.setText("已加载场景体系，可选正式场景，也可切换为自定义场景。")
        else:
            self.scene_hint_label.setText("未加载场景体系，请使用数据集自定义场景。")

    def on_scene_source_changed(self, source: str) -> None:
        is_taxonomy = source == "taxonomy"
        taxonomy_editable = is_taxonomy and bool(self.taxonomy.scenes)
        self.domain_combo.setEnabled(taxonomy_editable)
        self.scene_combo.setEnabled(taxonomy_editable)
        self.custom_en_edit.setEnabled(not is_taxonomy)
        self.custom_zh_edit.setEnabled(not is_taxonomy)
        self.mapping_edit.setEnabled(not is_taxonomy)
        if is_taxonomy and not self.taxonomy.scenes:
            self.scene_hint_label.setText("当前项目没有可用的场景体系；已有体系场景仅供查看。")

    def checked_quality_status(self) -> str:
        quality = "unreviewed"
        checked = self.quality_group.checkedButton()
        if checked:
            quality = checked.text()
        return quality

    def current_review_from_form(self) -> Review:
        quality = self.checked_quality_status()
        source = self.scene_source_combo.currentText()
        domain_id = self.domain_combo.currentData() or ""
        domain_text = self.domain_combo.currentText()
        level1_name = domain_text.split(" ", 1)[1].split(" / ")[0] if " " in domain_text else ""
        primary_scene = self.scene_combo.currentText() if source == "taxonomy" else ""
        custom_en = self.custom_en_edit.text().strip() if source != "taxonomy" else ""
        custom_zh = self.custom_zh_edit.text().strip() if source != "taxonomy" else ""
        scene_status = self.scene_status_combo.currentText()
        return Review(
            quality_status=quality,
            scene_status=scene_status,
            level1_id=domain_id if source == "taxonomy" else "",
            level1_name=level1_name if source == "taxonomy" else "",
            primary_level2_scene=primary_scene,
            scene_source=source,
            custom_scene_name_en=custom_en,
            custom_scene_name_zh=custom_zh,
            taxonomy_mapping_suggestion=self.mapping_edit.text().strip() if source != "taxonomy" else "",
            reviewer_note=self.note_edit.toPlainText(),
        )

    def save_quality_status_only(self) -> None:
        if self.loading or not self.db or self.dataset_id is None:
            return
        if self.current_index < 0 or self.current_index >= len(self.items):
            return
        sample, old = self.items[self.current_index]
        review = Review(
            quality_status=self.checked_quality_status(),
            scene_status=old.scene_status,
            level1_id=old.level1_id,
            level1_name=old.level1_name,
            primary_level2_scene=old.primary_level2_scene,
            secondary_scenes_json=old.secondary_scenes_json,
            scene_source=old.scene_source,
            custom_scene_name_en=old.custom_scene_name_en,
            custom_scene_name_zh=old.custom_scene_name_zh,
            taxonomy_mapping_suggestion=old.taxonomy_mapping_suggestion,
            issue_tags_json=old.issue_tags_json,
            reviewer_note=self.note_edit.toPlainText(),
        )
        self.db.save_review(sample.id, self.dataset_id, review)
        self.items[self.current_index] = (sample, review)
        self.update_sample_list_item(self.current_index, sample, review)
        self.status_label.setText(self.format_stats())

    def confirm_scene_selection(self) -> None:
        if self.current_index < 0 or self.current_index >= len(self.items):
            return
        source = self.scene_source_combo.currentText()
        if source == "taxonomy" and not self.scene_combo.currentText():
            QMessageBox.warning(self, "场景未选择", "请先选择一个正式二级场景。")
            return
        if source != "taxonomy" and not self.custom_en_edit.text().strip():
            QMessageBox.warning(self, "自定义场景未填写", "请填写自定义场景英文名。")
            return
        self.scene_status_combo.setCurrentText("assigned")
        self.save_current_review()
        self.scene_hint_label.setText("场景已确认并保存。")

    def reset_to_taxonomy_scene(self) -> None:
        if not self.taxonomy.scenes:
            QMessageBox.information(self, "未加载场景体系", "当前项目没有可用的场景体系 JSON，请使用自定义场景。")
            return
        self.scene_source_combo.setCurrentText("taxonomy")
        self.scene_combo.setCurrentIndex(0)
        self.custom_en_edit.clear()
        self.custom_zh_edit.clear()
        self.mapping_edit.clear()
        self.scene_status_combo.setCurrentText("unassigned")
        self.scene_hint_label.setText("已切换为正式体系选择，选好后点击确认场景。")

    def clear_scene_selection(self) -> None:
        self.scene_source_combo.setCurrentText(self.default_scene_source())
        self.scene_combo.setCurrentIndex(0)
        self.custom_en_edit.clear()
        self.custom_zh_edit.clear()
        self.mapping_edit.clear()
        self.scene_status_combo.setCurrentText("unassigned")
        self.save_current_review()
        self.scene_hint_label.setText("当前样本场景已清空。")

    def save_current_review(self) -> None:
        if self.loading or not self.db or self.dataset_id is None:
            return
        if self.current_index < 0 or self.current_index >= len(self.items):
            return
        sample, _old = self.items[self.current_index]
        review = self.current_review_from_form()
        self.db.save_review(sample.id, self.dataset_id, review)
        self.items[self.current_index] = (sample, review)
        self.update_sample_list_item(self.current_index, sample, review)
        self.update_common_scenes()
        self.status_label.setText(self.format_stats())

    def sample_list_item_text(self, index: int, sample: Sample, review: Review) -> str:
        scene = review.primary_level2_scene or review.custom_scene_name_en or "_"
        mark = {"accepted": "✓", "needs_correction": "!", "rejected": "x", "unreviewed": "·"}.get(review.quality_status, "·")
        return f"{index + 1}. {mark} {Path(sample.image_path).name}\n{scene} | {review.scene_status}"

    def update_sample_list_item(self, index: int, sample: Sample, review: Review) -> None:
        current = self.sample_list.item(index)
        if current:
            current.setText(self.sample_list_item_text(index, sample, review))

    def update_selected_export_button(self) -> None:
        if not hasattr(self, "selected_export_btn"):
            return
        count = len(self.sample_list.selectedIndexes())
        self.selected_export_btn.setText(f"导出选中（{count}）")
        self.selected_export_btn.setEnabled(count > 0)

    def update_common_scenes(self) -> None:
        if not self.db or self.dataset_id is None or not hasattr(self, "common_scene_list"):
            return
        self.common_scene_list.clear()
        for row in self.db.common_scenes(self.dataset_id):
            zh = f" / {row['scene_name_zh']}" if row["scene_name_zh"] else ""
            item = QListWidgetItem(f"{row['scene_name_en']}{zh}\n{row['scene_source']} | 已分配 {row['assigned_count']} 张")
            item.setData(Qt.ItemDataRole.UserRole, dict(row))
            self.common_scene_list.addItem(item)

    def apply_common_scene(self, item: QListWidgetItem) -> None:
        data = item.data(Qt.ItemDataRole.UserRole)
        if not isinstance(data, dict):
            return
        if data.get("scene_source") == "taxonomy" and not self.taxonomy.scenes:
            QMessageBox.information(
                self,
                "场景体系不可用",
                "该记录来自场景体系，但当前未加载对应 JSON。已有记录会保留，不能据此分配新样本。",
            )
            return
        self.scene_source_combo.setCurrentText(data.get("scene_source", "taxonomy"))
        if data.get("scene_source") == "taxonomy":
            idx = self.domain_combo.findData(data.get("level1_id", ""))
            if idx >= 0:
                self.domain_combo.setCurrentIndex(idx)
            self.reload_scene_combo()
            idx2 = self.scene_combo.findText(data.get("scene_name_en", ""))
            if idx2 >= 0:
                self.scene_combo.setCurrentIndex(idx2)
        else:
            self.custom_en_edit.setText(data.get("scene_name_en", ""))
            self.custom_zh_edit.setText(data.get("scene_name_zh", ""))
            self.mapping_edit.setText(data.get("mapped_level2_scene", ""))
        self.scene_status_combo.setCurrentText("assigned")
        self.scene_hint_label.setText("已填入常用场景候选，可继续修改，点击确认场景后保存。")

    def prev_sample(self) -> None:
        if not self.save_quality_note():
            return
        self.save_current_review()
        if self.current_index > 0:
            self.sample_list.setCurrentRow(self.current_index - 1)

    def next_sample(self) -> None:
        if not self.save_quality_note():
            return
        self.save_current_review()
        if self.current_index + 1 < len(self.items):
            self.sample_list.setCurrentRow(self.current_index + 1)

    def next_unreviewed(self) -> None:
        if not self.save_quality_note():
            return
        self.save_current_review()
        for idx in range(self.current_index + 1, len(self.items)):
            if self.items[idx][1].quality_status == "unreviewed" or self.items[idx][1].scene_status == "unassigned":
                self.sample_list.setCurrentRow(idx)
                return
        QMessageBox.information(self, "完成", "当前列表中没有后续未审或未分场景样本。")

    def format_stats(self) -> str:
        if not self.db or self.dataset_id is None:
            return ""
        stats = self.db.stats(self.dataset_id)
        quality = stats["quality"]
        reviewed = stats["total"] - quality.get("unreviewed", 0)
        return f"已审 {reviewed} / {stats['total']} | accepted {quality.get('accepted', 0)} | rejected {quality.get('rejected', 0)}"

    def default_export_directory(self) -> str:
        workspace_text = self.workspace_edit.text() if hasattr(self, "workspace_edit") else ""
        return resolve_export_start_directory(
            getattr(self, "last_export_directory", None),
            getattr(self, "workspace_dir", None),
            workspace_text,
        )

    def choose_export_directory(self, title: str) -> Path | None:
        selected = QFileDialog.getExistingDirectory(
            self,
            title,
            self.default_export_directory(),
            QFileDialog.Option.ShowDirsOnly,
        )
        if not selected:
            return None
        directory = Path(selected)
        self.last_export_directory = directory
        return directory

    def export_stats(self) -> None:
        if not self.db or self.dataset_id is None:
            return
        if not self.save_quality_note():
            return
        out = self.choose_export_directory("选择统计导出目录")
        if out is None:
            return
        stats = self.db.stats(self.dataset_id)
        summary_path = out / "review_summary.csv"
        scene_path = out / "scene_statistics.csv"
        with summary_path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["item", "count"])
            writer.writerow(["total", stats["total"]])
            for key, value in stats["quality"].items():
                writer.writerow([key, value])
            plan = build_export_plan(self.db.samples(self.dataset_id, "all", ""), True)
            writer.writerow(["accepted_with_assigned_scene", plan.assigned_total])
            writer.writerow(["accepted_with_unassigned_scene", plan.unassigned_total])
            writer.writerow(["accepted_with_uncertain_scene", plan.uncertain_total])
            writer.writerow(["accepted_with_invalid_assigned_scene", plan.invalid_assigned_total])
        with scene_path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["scene", "source", "count"])
            writer.writerows(stats["scenes"])
        QMessageBox.information(self, "导出完成", f"已导出：\n{summary_path}\n{scene_path}")

    def on_quality_note_changed(self) -> None:
        if self.loading:
            return
        self.note_save_label.setText("未保存")
        self.note_timer.start()

    def save_quality_note(self) -> bool:
        self.note_timer.stop()
        if self.loading or not self.db or self.dataset_id is None:
            return True
        if not 0 <= self.current_index < len(self.items):
            return True
        sample, old = self.items[self.current_index]
        text = self.note_edit.toPlainText()
        if text == old.reviewer_note:
            self.note_save_label.setText('已保存' if text else '')
            return True
        review = replace(old, reviewer_note=text)
        try:
            self.db.conn.execute(
                "UPDATE reviews SET reviewer_note = ?, updated_at = ? WHERE sample_id = ?",
                (text, now_text(), sample.id),
            )
            if self.db.conn.execute("SELECT changes()").fetchone()[0] == 0:
                self.db.save_review(sample.id, self.dataset_id, review)
            else:
                self.db.conn.commit()
        except sqlite3.Error as exc:
            self.db.conn.rollback()
            self.note_save_label.setText(f"保存失败：{exc}")
            return False
        self.items[self.current_index] = (sample, review)
        self.note_save_label.setText("已保存")
        return True

    def closeEvent(self, event) -> None:
        if self.save_quality_note():
            event.accept()
        else:
            event.ignore()

    def choose_quality_import_policy(
        self, plan: QualityImportPlan
    ) -> tuple[str, dict[str, str]] | None:
        counts = plan.counts()
        dialog = QDialog(self)
        dialog.setWindowTitle("预览质量进度导入")
        dialog.resize(1050, 620)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel(
            f"文件：{plan.source_path.name}\n"
            f"记录 {len(plan.entries)} 条 | 可导入 {counts['ready']} | "
            f"可补备注 {counts['note_fill']} | 相同 {counts['same']} | "
            f"冲突 {counts['conflict'] + counts['duplicate_conflict']} | "
            f"未匹配 {counts['unmatched']} | 无效 {counts['invalid']} | "
            f"重复 {counts['duplicate']}"
        ))
        layout.addWidget(QLabel("本功能只同步质量类型和质量备注，不会修改已有场景划分。"))

        policy_combo = QComboBox()
        policy_combo.addItem("仅填充本地未审核样本（推荐）", "fill_unreviewed")
        policy_combo.addItem("保留本地结果，并补充空白备注", "keep_local")
        policy_combo.addItem("采用导入结果，覆盖冲突样本", "use_imported")
        form = QFormLayout()
        form.addRow("合并策略", policy_combo)
        layout.addLayout(form)

        apply_label = QLabel()
        layout.addWidget(apply_label)

        noteworthy = [
            entry for entry in plan.entries
            if entry.category not in {"ready", "same"}
        ]
        shown = noteworthy[:500] if noteworthy else list(plan.entries[:100])
        table = QTableWidget(len(shown), 10)
        table.setHorizontalHeaderLabels([
            "Excel 行", "原图", "导入质量", "本地质量", "导入备注",
            "本地备注", "匹配方式", "处理状态", "本条处理", "说明"
        ])
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.verticalHeader().setVisible(False)
        conflict_choices: dict[str, QComboBox] = {}
        for row_index, entry in enumerate(shown):
            local_name = QUALITY_NAMES.get(entry.local_quality_status, "未审核" if entry.local_quality_status == "unreviewed" else "")
            values = (
                str(entry.source_row),
                entry.image_name,
                QUALITY_NAMES.get(entry.quality_status, entry.quality_status),
                local_name,
                entry.reviewer_note,
                entry.local_reviewer_note,
                entry.match_method,
                QUALITY_IMPORT_CATEGORY_NAMES.get(entry.category, entry.category),
                "",
                entry.detail,
            )
            for column, value in enumerate(values):
                item = QTableWidgetItem(value)
                if column in (1, 4, 5, 9):
                    item.setToolTip(value)
                table.setItem(row_index, column, item)
            if entry.category == "conflict":
                choice = QComboBox()
                choice.addItem("跟随批量策略", "")
                choice.addItem("保留本地", "keep_local")
                choice.addItem("采用导入", "use_imported")
                table.setCellWidget(row_index, 8, choice)
                conflict_choices[entry.matched_sample_id] = choice
            else:
                table.item(row_index, 8).setText("-")
        header = table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(9, QHeaderView.ResizeMode.Stretch)
        layout.addWidget(table, 1)
        if len(noteworthy) > len(shown):
            layout.addWidget(QLabel(f"需关注记录较多，表格仅显示前 {len(shown)} 条；全部记录仍会按所选策略处理。"))

        categories_by_policy = {
            "fill_unreviewed": {"ready"},
            "keep_local": {"ready", "note_fill"},
            "use_imported": {"ready", "note_fill", "conflict"},
        }

        def update_apply_count() -> None:
            policy = str(policy_combo.currentData())
            count = 0
            overwritten = 0
            for entry in plan.entries:
                apply_entry = entry.category in categories_by_policy[policy]
                if entry.category == "conflict" and entry.matched_sample_id in conflict_choices:
                    action = str(conflict_choices[entry.matched_sample_id].currentData())
                    if action:
                        apply_entry = action == "use_imported"
                count += apply_entry
                overwritten += entry.category == "conflict" and apply_entry
            conflict_text = f"；其中覆盖冲突 {overwritten} 条" if overwritten else ""
            apply_label.setText(
                f"确认后将写入 {count} 条记录{conflict_text}。"
                "未匹配、无效和表内冲突不会写入。"
            )

        policy_combo.currentIndexChanged.connect(update_apply_count)
        for choice in conflict_choices.values():
            choice.currentIndexChanged.connect(update_apply_count)
        update_apply_count()
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("确认合并")
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return None
        overrides = {
            sample_id: str(choice.currentData())
            for sample_id, choice in conflict_choices.items()
            if choice.currentData()
        }
        return str(policy_combo.currentData()), overrides

    def import_quality_progress(self) -> None:
        if not self.db or self.dataset_id is None or not self.save_quality_note():
            return
        paths, _selected_filter = QFileDialog.getOpenFileNames(
            self,
            "选择成员导出的质量审核记录",
            self.default_export_directory(),
            "质量审核记录 (*.xlsx)",
        )
        if not paths:
            return
        dataset = self.db.conn.execute(
            "SELECT image_root, mask_root FROM datasets WHERE id = ?",
            (self.dataset_id,),
        ).fetchone()
        if dataset is None:
            QMessageBox.critical(self, "导入失败", "当前数据集记录不存在。")
            return

        reports: list[str] = []
        for path_text in paths:
            path = Path(path_text)
            try:
                items = self.db.samples(self.dataset_id, "all", "")
                plan = read_quality_progress_xlsx(
                    path,
                    items,
                    Path(dataset["image_root"]),
                    Path(dataset["mask_root"]),
                )
            except Exception as exc:
                QMessageBox.critical(self, "读取失败", f"{path.name} 无法读取：\n{exc}")
                continue

            if self.db.quality_import_seen(self.dataset_id, plan.source_sha256):
                answer = QMessageBox.question(
                    self,
                    "文件已经导入过",
                    f"{path.name} 的内容此前已经导入过。\n"
                    "再次导入通常不会增加进度，是否仍要查看并继续？",
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                    QMessageBox.StandardButton.No,
                )
                if answer != QMessageBox.StandardButton.Yes:
                    reports.append(f"{path.name}：已导入过，本次跳过")
                    continue

            selection = self.choose_quality_import_policy(plan)
            if selection is None:
                reports.append(f"{path.name}：已取消")
                continue
            policy, conflict_overrides = selection
            categories_by_policy = {
                "fill_unreviewed": {"ready"},
                "keep_local": {"ready", "note_fill"},
                "use_imported": {"ready", "note_fill", "conflict"},
            }
            expected = 0
            for entry in plan.entries:
                apply_entry = entry.category in categories_by_policy[policy]
                if entry.category == "conflict":
                    action = conflict_overrides.get(entry.matched_sample_id)
                    if action:
                        apply_entry = action == "use_imported"
                expected += apply_entry
            backup_path: Path | None = None
            try:
                if expected:
                    backup_root = self.db.db_path.parent / "backups"
                    backup_path = backup_root / (
                        f"review_before_quality_import_{time.strftime('%Y%m%d_%H%M%S')}_"
                        f"{time.time_ns() % 1000000:06d}.sqlite3"
                    )
                    self.db.backup_to(backup_path)
                _import_id, applied = self.db.apply_quality_import(
                    self.dataset_id, plan, policy, conflict_overrides
                )
            except Exception as exc:
                QMessageBox.critical(
                    self,
                    "合并失败",
                    f"{path.name} 未写入工作区：\n{exc}"
                    + (f"\n合并前备份：{backup_path}" if backup_path else ""),
                )
                continue
            counts = plan.counts()
            overwritten = sum(
                entry.category == "conflict"
                and (
                    conflict_overrides.get(entry.matched_sample_id) == "use_imported"
                    or (
                        entry.matched_sample_id not in conflict_overrides
                        and policy == "use_imported"
                    )
                )
                for entry in plan.entries
            )
            report = (
                f"{path.name}：写入 {applied}，覆盖冲突 {overwritten}，"
                f"冲突跳过 {counts['conflict'] - overwritten}，"
                f"未匹配 {counts['unmatched']}，无效 {counts['invalid']}"
            )
            if backup_path:
                report += f"；备份 {backup_path.name}"
            reports.append(report)

        if reports:
            self.reload_samples()
            QMessageBox.information(self, "质量进度导入结果", "\n".join(reports))

    def choose_quality_export_options(
        self,
        items: list[tuple[Sample, Review]],
        title: str,
        scope_text: str = "",
    ) -> tuple[set[str], bool] | None:
        counts = Counter(review.quality_status for _, review in items)
        dialog = QDialog(self)
        dialog.setWindowTitle(title)
        layout = QVBoxLayout(dialog)
        if scope_text:
            layout.addWidget(QLabel(scope_text))
        table_checkbox = QCheckBox(f"质量审核记录.xlsx（全部已审核 {sum(counts[status] for status in QUALITY_NAMES)} 条）")
        table_checkbox.setChecked(True)
        layout.addWidget(table_checkbox)
        choices: dict[str, QCheckBox] = {}
        for status, name in QUALITY_NAMES.items():
            label = "需修改（待确认）" if status == "needs_correction" else name
            checkbox = QCheckBox(f"{label}图片与掩膜（{counts[status]} 张）")
            checkbox.setChecked(True)
            layout.addWidget(checkbox)
            choices[status] = checkbox
        layout.addWidget(QLabel(f"未审核 {counts['unreviewed']} 张，本次不导出"))
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return None
        statuses = {status for status, checkbox in choices.items() if checkbox.isChecked()}
        if not table_checkbox.isChecked() and not statuses:
            QMessageBox.information(self, "没有选择导出项", "请至少勾选表格或一个质量类别。")
            return None
        if not any(counts[status] for status in (set(QUALITY_NAMES) if table_checkbox.isChecked() else statuses)):
            QMessageBox.information(self, "没有可导出的样本", "所选项目中没有已审核样本。")
            return None
        return statuses, table_checkbox.isChecked()

    def export_quality(self) -> None:
        if not self.db or self.dataset_id is None or not self.save_quality_note():
            return
        items = self.db.samples(self.dataset_id, "all", "")
        options = self.choose_quality_export_options(items, "按质量导出")
        if options is None:
            return
        file_statuses, include_table = options
        destination = self.choose_export_directory("选择质量导出目录")
        if destination is None:
            return
        row = self.db.conn.execute("SELECT image_root, mask_root FROM datasets WHERE id = ?", (self.dataset_id,)).fetchone()
        root = destination / f"quality_export_{time.strftime('%Y%m%d_%H%M%S')}_{time.time_ns() % 1000000:06d}"
        try:
            count = export_quality_items(
                items, root, Path(row["image_root"]), Path(row["mask_root"]),
                file_statuses, True, include_table,
                set(QUALITY_NAMES) if include_table else None,
            )
        except Exception as exc:
            QMessageBox.critical(self, "导出失败", f"质量导出未完成：\n{exc}")
            return
        QMessageBox.information(self, "导出完成", f"已导出 {count} 条质量审核记录：\n{root}")

    def export_selected(self) -> None:
        if not self.db or self.dataset_id is None or not self.save_quality_note():
            return
        selected_ids = [
            str(index.data(Qt.ItemDataRole.UserRole))
            for index in sorted(self.sample_list.selectedIndexes(), key=lambda index: index.row())
            if index.data(Qt.ItemDataRole.UserRole)
        ]
        if not selected_ids:
            QMessageBox.information(self, "没有选中样本", "请先在左侧列表选择要导出的样本。")
            return

        all_items = self.db.samples(self.dataset_id, "all", "")
        items_by_id = {sample.id: (sample, review) for sample, review in all_items}
        selected_items = [items_by_id[sample_id] for sample_id in selected_ids if sample_id in items_by_id]
        if len(selected_items) != len(selected_ids):
            QMessageBox.warning(self, "选中样本已变化", "部分选中样本已不存在，请刷新列表后重试。")
            return

        options = self.choose_quality_export_options(
            selected_items,
            "按质量导出选中样本",
            f"当前选中 {len(selected_items)} 张；只统计和导出这些样本。",
        )
        if options is None:
            return
        file_statuses, include_table = options
        destination = self.choose_export_directory("选择选中样本导出目录")
        if destination is None:
            return
        dataset = self.db.conn.execute(
            "SELECT image_root, mask_root FROM datasets WHERE id = ?", (self.dataset_id,)
        ).fetchone()
        if dataset is None:
            QMessageBox.critical(self, "导出失败", "当前数据集记录不存在。")
            return
        root = destination / (
            f"selected_quality_export_{time.strftime('%Y%m%d_%H%M%S')}_"
            f"{time.time_ns() % 1000000:06d}"
        )
        try:
            exported = export_quality_items(
                selected_items,
                root,
                Path(dataset["image_root"]),
                Path(dataset["mask_root"]),
                file_statuses,
                True,
                include_table,
                set(QUALITY_NAMES) if include_table else None,
            )
        except Exception as exc:
            QMessageBox.critical(self, "导出失败", f"选中样本导出未完成：\n{exc}")
            return
        QMessageBox.information(
            self, "导出完成", f"已按质量导出 {exported} 个选中样本：\n{root}"
        )

    def export_accepted(self) -> None:
        self._export_review_data(group_by_scene=False)

    def export_by_scene(self) -> None:
        self._export_review_data(group_by_scene=True)

    def _export_review_data(self, group_by_scene: bool) -> None:
        if not self.db or self.dataset_id is None:
            return
        if not self.save_quality_note():
            return
        self.save_current_review()
        plan = build_export_plan(self.db.samples(self.dataset_id, "all", ""), group_by_scene)
        if not plan.items:
            detail = "当前没有质量状态为 accepted 的样本。"
            if group_by_scene and plan.accepted_total:
                detail = "存在 accepted 样本，但没有同时满足“场景已确认且场景名有效”的样本。"
            QMessageBox.information(self, "没有可导出的样本", detail)
            return
        if group_by_scene:
            prompt = (
                f"质量合格：{plan.accepted_total} 张\n"
                f"场景已确认：{plan.assigned_total} 张\n"
                f"场景未分配：{plan.unassigned_total} 张\n"
                f"场景不确定：{plan.uncertain_total} 张\n"
                f"场景状态异常：{plan.invalid_assigned_total} 张\n\n"
                f"本次将按场景导出 {len(plan.items)} 张，是否继续？"
            )
            dialog_title = "确认按场景导出"
            directory_title = "选择场景导出目录"
            directory_prefix = "export_by_scene"
        else:
            prompt = f"本次将导出全部 {plan.accepted_total} 张质量合格样本，不限制场景状态。是否继续？"
            dialog_title = "确认导出合格数据"
            directory_title = "选择合格数据导出目录"
            directory_prefix = "export_accepted"
        answer = QMessageBox.question(
            self,
            dialog_title,
            prompt,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.Yes,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        out_dir = self.choose_export_directory(directory_title)
        if out_dir is None:
            return
        root = out_dir / f"{directory_prefix}_{time.strftime('%Y%m%d_%H%M%S')}"
        try:
            row = self.db.conn.execute("SELECT image_root, mask_root FROM datasets WHERE id = ?", (self.dataset_id,)).fetchone()
            exported = export_review_items(plan, root, group_by_scene, Path(row["image_root"]), Path(row["mask_root"]))
        except (OSError, ValueError) as exc:
            QMessageBox.critical(self, "导出失败", f"导出过程中发生文件错误：\n{exc}")
            return
        description = "质量合格样本" if not group_by_scene else "质量合格且场景已确认样本"
        QMessageBox.information(self, "导出完成", f"已导出 {exported} 个{description}到：\n{root}")

    def keyPressEvent(self, event) -> None:  # type: ignore[no-untyped-def]
        key = event.key()
        if key in (Qt.Key.Key_Right, Qt.Key.Key_D):
            self.next_sample()
        elif key in (Qt.Key.Key_Left, Qt.Key.Key_A):
            self.prev_sample()
        elif key == Qt.Key.Key_Space:
            self.mask_checkbox.setChecked(not self.mask_checkbox.isChecked())
        elif key == Qt.Key.Key_F:
            self.image_canvas.fit()
        elif key == Qt.Key.Key_1:
            self.set_quality("accepted")
        elif key == Qt.Key.Key_2:
            self.set_quality("needs_correction")
        elif key == Qt.Key.Key_3:
            self.set_quality("rejected")
        else:
            super().keyPressEvent(event)

    def set_quality(self, status: str) -> None:
        for button in self.quality_group.buttons():
            if button.text() == status:
                button.setChecked(True)
                self.save_quality_status_only()
                return


def main() -> None:
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
